from copy import deepcopy
from io import BytesIO

import openpyxl
import pytest

import app as application


def set_session(client, user_id=7):
    with client.session_transaction() as session:
        session["uid"] = user_id
        session["username"] = "Excel Tester"


def report_data():
    return {
        "date_range": {"start_date": "2026-05-01", "end_date": "2026-05-31"},
        "expenses": {
            "expense_summary": {
                "total_expenses": 600.5,
                "transaction_count": 3,
                "average_expense": 200.1667,
                "largest_expense": 300,
            },
            "category_totals": [{"category": "Food", "amount": 600.5}],
            "monthly_totals": [{"month": "2026-05", "amount": 600.5}],
        },
        "budgets": {
            "total_budget": 1000,
            "total_spent": 750,
            "total_remaining": 250,
            "overall_utilization": 75,
            "budget_count": 1,
            "budgets": [{
                "budget_name": "Food Budget",
                "category": "Food",
                "budget_amount": 1000,
                "spent_amount": 750,
                "remaining_amount": 250,
                "status": "Active",
            }],
        },
        "investments": {
            "holding_count": 1,
            "holdings": [{
                "asset_name": "Index Fund",
                "asset_type": "Stocks",
                "quantity": 2,
                "purchase_price": 500,
                "current_value": 1200,
                "invested_value": 1000,
                "purchase_date": "2026-01-10",
            }],
            "stats": {
                "total_invested": 1000,
                "current_value": 1200,
                "absolute_return": 200,
                "return_percentage": 20,
                "allocation": [{"asset_type": "Stocks", "percentage": 100}],
            },
        },
        "goals": {
            "goal_count": 1,
            "goals": [{
                "goal_name": "Emergency Fund",
                "goal_category": "Savings",
                "target_amount": 1000,
                "current_amount": 250,
                "progress_percentage": 25,
                "remaining_amount": 750,
                "status": "Active",
            }],
        },
        "financial_health": {
            "available": True,
            "score": 79,
            "grade": "Good",
            "coverage": 80,
            "earned_points": 63,
            "available_weight": 80,
            "components": {
                "budget": {
                    "available": True,
                    "score": 35,
                    "message": "Budget usage is healthy.",
                },
                "spending": {
                    "available": False,
                    "score": 0,
                    "message": "Expense data is unavailable.",
                },
            },
        },
        "unavailable_metrics": {
            "total_income": {
                "available": False,
                "value": None,
                "reason": "No reliable actual income source exists.",
            },
            "total_savings": {
                "available": False,
                "value": None,
                "reason": "No actual savings balance exists.",
            },
            "recent_alerts": {
                "available": False,
                "value": None,
                "reason": "No persisted notification or alert data source exists.",
            },
        },
    }


def load_response_workbook(response):
    return openpyxl.load_workbook(BytesIO(response.data), read_only=True, data_only=False)


def test_authenticated_user_can_export_valid_excel_workbook(monkeypatch):
    data = report_data()
    monkeypatch.setattr(application, "build_reporting_data", lambda *args, **kwargs: deepcopy(data))
    client = application.app.test_client()
    set_session(client)

    response = client.get("/reports/export/excel")
    workbook = load_response_workbook(response)

    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response.headers["Content-Disposition"] == (
        "attachment; filename=\"finsight_financial_report.xlsx\""
    )
    assert workbook.sheetnames == [
        "Summary",
        "Expense Analysis",
        "Budget Summary",
        "Investments",
        "Goals",
        "Financial Health",
    ]
    workbook.close()


def test_workbook_contains_report_values_and_current_snapshots(monkeypatch):
    data = report_data()
    monkeypatch.setattr(application, "build_reporting_data", lambda *args, **kwargs: deepcopy(data))
    client = application.app.test_client()
    set_session(client)

    workbook = load_response_workbook(client.get("/reports/export/excel"))

    assert workbook["Summary"]["A8"].value == 600.5
    assert workbook["Summary"]["B8"].value == 3
    assert workbook["Summary"]["A12"].value == 79
    assert workbook["Expense Analysis"]["A12"].value == "Food"
    assert workbook["Expense Analysis"]["B12"].value == 600.5
    assert workbook["Budget Summary"]["A14"].value == "Food Budget"
    assert workbook["Budget Summary"]["C14"].value == 1000
    assert workbook["Investments"]["A14"].value == "Index Fund"
    assert workbook["Goals"]["A8"].value == "Emergency Fund"
    assert workbook["Financial Health"]["A8"].value == 79
    assert "No reliable actual income source exists." in [
        cell.value
        for row in workbook["Summary"].iter_rows()
        for cell in row
        if cell.value
    ]
    workbook.close()


def test_excel_export_requires_authentication():
    response = application.app.test_client().get("/reports/export/excel")

    assert response.status_code == 302
    assert response.headers["Location"].endswith("/login")


def test_excel_export_uses_session_user_and_preserves_date_range(monkeypatch):
    calls = []

    def fake_build(user_id, start_date, end_date, *, investment_service, goal_service):
        calls.append((user_id, start_date, end_date))
        return report_data()

    monkeypatch.setattr(application, "build_reporting_data", fake_build)
    client = application.app.test_client()
    set_session(client, user_id=7)

    response = client.get(
        "/reports/export/excel?user_id=99&start_date=2026-01-01&end_date=2026-03-31"
    )

    assert response.status_code == 200
    assert calls == [(7, "2026-01-01", "2026-03-31")]


@pytest.mark.parametrize(
    "query_string",
    [
        {"start_date": "2026-01-01"},
        {"start_date": "invalid", "end_date": "2026-03-31"},
        {"start_date": "2026-04-01", "end_date": "2026-03-31"},
    ],
)
def test_excel_export_rejects_invalid_date_ranges(query_string):
    client = application.app.test_client()
    set_session(client)

    response = client.get("/reports/export/excel", query_string=query_string)

    assert response.status_code == 400
    assert response.is_json
    assert response.get_json()["error"]


def test_empty_expenses_still_generate_valid_workbook(monkeypatch):
    data = report_data()
    data["date_range"] = {"start_date": None, "end_date": None}
    data["expenses"]["expense_summary"] = {
        "total_expenses": 0,
        "transaction_count": 0,
        "average_expense": 0,
        "largest_expense": 0,
    }
    data["expenses"]["category_totals"] = []
    data["expenses"]["monthly_totals"] = []
    monkeypatch.setattr(application, "build_reporting_data", lambda *args, **kwargs: data)
    client = application.app.test_client()
    set_session(client)

    workbook = load_response_workbook(client.get("/reports/export/excel"))
    values = [
        cell.value
        for row in workbook["Expense Analysis"].iter_rows()
        for cell in row
        if cell.value
    ]

    assert "All Available Expense History" in values
    assert "No expense data available for the selected period." in values
    assert "No monthly expense data available." in values
    workbook.close()


def test_excel_export_does_not_mutate_report_data(monkeypatch):
    data = report_data()
    original = deepcopy(data)
    monkeypatch.setattr(application, "build_reporting_data", lambda *args, **kwargs: data)
    client = application.app.test_client()
    set_session(client)

    response = client.get("/reports/export/excel")

    assert response.status_code == 200
    assert data == original


def test_reports_page_enables_excel_link_with_selected_dates(monkeypatch):
    monkeypatch.setattr(application, "build_reporting_data", lambda *args, **kwargs: report_data())
    client = application.app.test_client()
    set_session(client)

    response = client.get(
        "/reports?start_date=2026-05-01&end_date=2026-05-31",
        headers={"Accept": "text/html"},
    )

    assert response.status_code == 200
    assert (
        'href="/reports/export/excel?start_date=2026-05-01&amp;end_date=2026-05-31"'
        in response.get_data(as_text=True)
    )
