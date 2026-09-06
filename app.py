import csv
import base64
import hashlib
import os
import re
import secrets
import time
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from html import escape
from io import BytesIO, StringIO

from flask import (
    Flask,
    flash,
    g,
    jsonify,
    redirect,
    render_template,
    Response,
    request,
    session,
    url_for,
)
import pyotp
from cryptography.fernet import Fernet, InvalidToken
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from werkzeug.security import generate_password_hash
from werkzeug.middleware.proxy_fix import ProxyFix

from config import Config
from db import (
    create_budget,
    create_user_session,
    delete_budget,
    disable_totp_for_user,
    enable_totp_for_user,
    filter_budgets,
    get_budget,
    get_monthly_expense_summary,
    get_notifications,
    get_unread_notification_count,
    get_expense_summary,
    get_summary_stats,
    get_transactions,
    get_totp_credential,
    get_totp_status,
    get_user_by_id,
    get_user_preferences,
    is_user_session_active,
    ensure_user_preferences,
    init_db,
    list_preference_currencies,
    list_preference_languages,
    list_active_user_sessions,
    login_user,
    mark_all_notifications_read,
    mark_notification_read,
    register_user,
    create_remember_me_token,
    rotate_remember_me_token,
    revoke_remember_me_token,
    revoke_all_remember_me_tokens,
    create_pending_registration,
    get_pending_registration,
    replace_pending_registration_otp,
    complete_pending_registration,
    get_registration_resend_status,
    registration_otp_digest,
    get_user_for_password_reset,
    get_password_reset_challenge,
    get_password_reset_challenge_by_id,
    create_password_reset_challenge,
    get_password_reset_resend_status,
    replace_password_reset_otp,
    verify_password_reset_otp,
    get_password_reset_authorization,
    reset_password_with_authorization,
    revoke_all_user_sessions,
    revoke_current_user_session,
    revoke_user_session,
    save_pending_totp_secret,
    update_user_preferences,
    update_user_password_and_revoke_other_sessions,
    verify_user_password,
    update_budget,
)
from investment_repository import InvestmentRepository
from investment_service import ASSET_TYPES, InvestmentService
from goal_repository import GoalRepository
from goal_service import GOAL_CATEGORIES, GoalService
from financial_health_service import calculate_financial_health
from i18n import TRANSLATIONS, currency_symbol, format_currency, translate
from report_service import ReportValidationError, build_reporting_data
from email_service import email_service

app = Flask(__name__)


def _is_production_environment():
    production_values = {"production", "prod"}
    environment_names = ("FLASK_ENV", "APP_ENV", "ENVIRONMENT", "VERCEL_ENV", "ENV")
    return any(
        os.getenv(name, "").strip().lower() in production_values
        for name in environment_names
    )


def _get_secret_key():
    if Config.SECRET_KEY:
        return Config.SECRET_KEY
    if _is_production_environment():
        raise RuntimeError("SECRET_KEY must be set in production.")
    return secrets.token_urlsafe(32)


def _debug_enabled():
    if _is_production_environment():
        return False
    return os.getenv("FLASK_DEBUG", "").strip().lower() in {
        "1", "true", "yes", "on"
    }


if _is_production_environment():
    app.wsgi_app = ProxyFix(app.wsgi_app, x_proto=1, x_host=1)


app.config["SECRET_KEY"] = _get_secret_key()


DEFAULT_SESSION_INACTIVITY_TIMEOUT_SECONDS = 30 * 60
DEFAULT_SESSION_ABSOLUTE_TIMEOUT_SECONDS = 24 * 60 * 60
DEFAULT_REMEMBER_ME_TIMEOUT_SECONDS = 30 * 24 * 60 * 60
REMEMBER_ME_COOKIE_NAME = "finsight_remember_me"
PASSWORD_RESET_OTP_TIMEOUT_SECONDS = 10 * 60
PASSWORD_RESET_AUTH_TIMEOUT_SECONDS = 10 * 60
PASSWORD_RESET_GENERIC_MESSAGE = (
    "If the account exists, a verification code has been sent to the email address."
)
PASSWORD_RESET_INVALID_MESSAGE = "Invalid or expired verification code."


def _session_timeout_setting(name, default):
    raw_value = os.getenv(name)
    if raw_value is None or not raw_value.strip():
        return default
    try:
        value = int(raw_value)
    except (TypeError, ValueError) as exc:
        raise RuntimeError(f"{name} must be a positive integer.") from exc
    if value <= 0:
        raise RuntimeError(f"{name} must be a positive integer.")
    return value


app.config.update(
    SESSION_INACTIVITY_TIMEOUT_SECONDS=_session_timeout_setting(
        "SESSION_INACTIVITY_TIMEOUT_SECONDS",
        DEFAULT_SESSION_INACTIVITY_TIMEOUT_SECONDS,
    ),
    SESSION_ABSOLUTE_TIMEOUT_SECONDS=_session_timeout_setting(
        "SESSION_ABSOLUTE_TIMEOUT_SECONDS",
        DEFAULT_SESSION_ABSOLUTE_TIMEOUT_SECONDS,
    ),
    REMEMBER_ME_TIMEOUT_SECONDS=_session_timeout_setting(
        "REMEMBER_ME_TIMEOUT_SECONDS",
        DEFAULT_REMEMBER_ME_TIMEOUT_SECONDS,
    ),
)
if app.config["REMEMBER_ME_TIMEOUT_SECONDS"] <= app.config["SESSION_ABSOLUTE_TIMEOUT_SECONDS"]:
    raise RuntimeError(
        "REMEMBER_ME_TIMEOUT_SECONDS must exceed SESSION_ABSOLUTE_TIMEOUT_SECONDS."
    )


def _secure_session_cookie_enabled():
    if _is_production_environment():
        return True
    return os.getenv("SESSION_COOKIE_SECURE", "").strip().lower() in {
        "1", "true", "yes", "on"
    }


def _configure_security_settings():
    app.config.update(
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_SECURE=_secure_session_cookie_enabled(),
    )


SECURITY_HEADERS = {
    "X-Content-Type-Options": "nosniff",
    "X-Frame-Options": "SAMEORIGIN",
    "Referrer-Policy": "strict-origin-when-cross-origin",
    "Permissions-Policy": "camera=(), geolocation=(), microphone=()",
    "Content-Security-Policy": (
        "default-src 'self'; "
        "base-uri 'self'; "
        "object-src 'none'; "
        "frame-ancestors 'self'; "
        "form-action 'self'; "
        "script-src 'self' 'unsafe-inline' https://cdn.jsdelivr.net; "
        "style-src 'self' 'unsafe-inline' https://fonts.googleapis.com "
        "https://cdnjs.cloudflare.com https://cdn.jsdelivr.net; "
        "font-src 'self' https://fonts.gstatic.com https://cdnjs.cloudflare.com; "
        "img-src 'self' data:; "
        "connect-src 'self'"
    ),
}


@app.after_request
def add_security_headers(response):
    for header, value in SECURITY_HEADERS.items():
        response.headers.setdefault(header, value)
    if _is_production_environment() and request.is_secure:
        response.headers.setdefault(
            "Strict-Transport-Security", "max-age=31536000"
        )
    return response


_configure_security_settings()


LOGIN_MAX_FAILURES = 5
LOGIN_RATE_WINDOW_SECONDS = 300
_login_failures = {}


def _login_rate_limit_key():
    return request.remote_addr or "unknown"


def _active_login_failures(key, now=None):
    now = time.monotonic() if now is None else now
    timestamps = [
        timestamp
        for timestamp in _login_failures.get(key, [])
        if now - timestamp < LOGIN_RATE_WINDOW_SECONDS
    ]
    if timestamps:
        _login_failures[key] = timestamps
    else:
        _login_failures.pop(key, None)
    return timestamps


def _login_rate_limited():
    return len(_active_login_failures(_login_rate_limit_key())) >= LOGIN_MAX_FAILURES


def _record_login_failure():
    key = _login_rate_limit_key()
    timestamps = _active_login_failures(key)
    timestamps.append(time.monotonic())
    _login_failures[key] = timestamps


def _clear_login_failures():
    _login_failures.pop(_login_rate_limit_key(), None)

init_db()
investment_service = InvestmentService(InvestmentRepository())
goal_service = GoalService(GoalRepository())


BUDGET_HEALTHY_LIMIT = Decimal("80")
BUDGET_APPROACHING_LIMIT = Decimal("100")
CATEGORY_CONCENTRATION_THRESHOLD = Decimal("0.50")
LARGE_EXPENSE_MULTIPLIER = Decimal("2")
APPROACHING_BUDGET_THRESHOLD = Decimal("80")
OVER_BUDGET_THRESHOLD = Decimal("100")
PREFERENCE_THEME_OPTIONS = ("default", "light", "dark")
PREFERENCE_THEME_DISPLAY_OPTIONS = ("light", "dark")
PREFERENCE_BOOLEAN_FIELDS = (
    "budget_overspending_alerts",
    "weekly_savings_digest_enabled",
    "sip_due_date_reminders_enabled",
    "bill_due_date_reminders_enabled",
)
PREFERENCE_DEFAULTS = {
    "theme": "default",
    "currency": "USD",
    "language": "en",
    "budget_overspending_alerts": False,
    "weekly_savings_digest_enabled": False,
    "sip_due_date_reminders_enabled": False,
    "bill_due_date_reminders_enabled": False,
}


def normalize_theme(theme):
    """Map persisted theme values to the themes supported by the UI."""
    return "dark" if str(theme or "").strip().lower() == "dark" else "light"


def sync_theme_session(preferences):
    session["theme"] = normalize_theme((preferences or {}).get("theme"))


def login_required_redirect():
    user_id = session.get("uid")
    session_token_hash = current_session_token_hash()
    if type(user_id) is not int or user_id <= 0 or not session_token_hash:
        session.clear()
        return redirect(url_for("login"))
    if not is_user_session_active(
        user_id,
        session_token_hash,
        app.config["SESSION_INACTIVITY_TIMEOUT_SECONDS"],
        app.config["SESSION_ABSOLUTE_TIMEOUT_SECONDS"],
    ):
        session.clear()
        flash("Your session is no longer active. Please sign in again.", "danger")
        return redirect(url_for("login"))

    return None


def current_user_id():
    return session["uid"]


def current_session_token_hash():
    session_token = session.get("auth_session_token")
    if not isinstance(session_token, str) or not session_token:
        return None
    return hashlib.sha256(session_token.encode("utf-8")).hexdigest()


def request_device_info():
    return (request.headers.get("User-Agent") or "Unknown device").strip()[:255]


def request_ip_address():
    remote_addr = request.remote_addr or ""
    return remote_addr[:45] or None


def _totp_cipher():
    key_material = hashlib.sha256(
        app.config["SECRET_KEY"].encode("utf-8")
    ).digest()
    return Fernet(base64.urlsafe_b64encode(key_material))


def encrypt_totp_secret(secret):
    return _totp_cipher().encrypt(secret.encode("utf-8")).decode("utf-8")


def decrypt_totp_secret(secret_encrypted):
    return _totp_cipher().decrypt(secret_encrypted.encode("utf-8")).decode("utf-8")


def establish_authenticated_session(user):
    """Create the Flask and tracked-device session after all auth checks pass."""
    session_token = secrets.token_urlsafe(32)
    try:
        create_user_session(
            user["id"],
            hashlib.sha256(session_token.encode("utf-8")).hexdigest(),
            request_device_info(),
            request_ip_address(),
        )
    except Exception:
        app.logger.exception("Authenticated session registration failed")
        return False

    session.clear()
    session["uid"] = user["id"]
    session["username"] = user["username"]
    session["email"] = user["email"]
    session["auth_session_token"] = session_token
    try:
        preferences = get_user_preferences(user["id"])
    except Exception:
        app.logger.exception("Unable to load the user's theme preference")
        preferences = None
    sync_theme_session(preferences)
    return True


def _remember_me_token_hash(token):
    return hashlib.sha256(token.encode("utf-8")).hexdigest()


def _remember_me_requested():
    return request.form.get("remember_me", "").strip().lower() in {
        "1", "true", "on", "yes"
    }


def _issue_remember_me_token(user_id):
    raw_token = secrets.token_urlsafe(32)
    try:
        create_remember_me_token(
            user_id,
            _remember_me_token_hash(raw_token),
            app.config["REMEMBER_ME_TIMEOUT_SECONDS"],
            request_device_info(),
            request_ip_address(),
        )
    except Exception:
        app.logger.exception("Remember Me credential creation failed")
        return None
    return raw_token


def _set_remember_me_cookie(response, raw_token):
    g.clear_remember_me_cookie = False
    timeout = app.config["REMEMBER_ME_TIMEOUT_SECONDS"]
    response.set_cookie(
        REMEMBER_ME_COOKIE_NAME,
        raw_token,
        max_age=timeout,
        expires=datetime.now(timezone.utc) + timedelta(seconds=timeout),
        path="/",
        httponly=True,
        secure=app.config["SESSION_COOKIE_SECURE"],
        samesite=app.config["SESSION_COOKIE_SAMESITE"],
    )
    return response


def _clear_remember_me_cookie(response):
    g.remember_me_cookie_token = None
    g.clear_remember_me_cookie = False
    response.delete_cookie(
        REMEMBER_ME_COOKIE_NAME,
        path="/",
        secure=app.config["SESSION_COOKIE_SECURE"],
        httponly=True,
        samesite=app.config["SESSION_COOKIE_SAMESITE"],
    )
    return response


def _forget_remember_me_cookie():
    g.remember_me_cookie_token = None
    g.clear_remember_me_cookie = True


def _revoke_remember_me_token(user_id, raw_token):
    if not raw_token:
        return
    try:
        revoke_remember_me_token(user_id, _remember_me_token_hash(raw_token))
    except Exception:
        app.logger.exception("Remember Me credential revocation failed")


@app.after_request
def apply_remember_me_cookie(response):
    raw_token = getattr(g, "remember_me_cookie_token", None)
    if raw_token:
        _set_remember_me_cookie(response, raw_token)
    elif getattr(g, "clear_remember_me_cookie", False):
        _clear_remember_me_cookie(response)
    return response


@app.before_request
def restore_remembered_session():
    if (
        request.endpoint == "static"
        or session.get("uid")
        or session.get("pending_2fa_user_id")
    ):
        return None

    raw_token = request.cookies.get(REMEMBER_ME_COOKIE_NAME)
    if not raw_token:
        return None

    new_raw_token = secrets.token_urlsafe(32)
    try:
        user_id = rotate_remember_me_token(
            _remember_me_token_hash(raw_token),
            _remember_me_token_hash(new_raw_token),
            app.config["REMEMBER_ME_TIMEOUT_SECONDS"],
            request_device_info(),
            request_ip_address(),
        )
    except Exception:
        app.logger.exception("Remember Me credential validation failed")
        g.clear_remember_me_cookie = True
        return None

    if not user_id:
        g.clear_remember_me_cookie = True
        return None

    try:
        user = get_user_by_id(user_id)
        if not user:
            g.clear_remember_me_cookie = True
            _revoke_remember_me_token(user_id, new_raw_token)
            return None
        totp_status = get_totp_status(user_id)
    except Exception:
        app.logger.exception("Remember Me user validation failed")
        g.clear_remember_me_cookie = True
        _revoke_remember_me_token(user_id, new_raw_token)
        return None

    g.remember_me_cookie_token = new_raw_token
    if (totp_status or {}).get("is_enabled"):
        session.clear()
        session["pending_2fa_user_id"] = user_id
        session["pending_2fa_remembered_session"] = True
        return redirect(url_for("login_two_factor"))

    if not establish_authenticated_session(user):
        session.clear()
        g.clear_remember_me_cookie = True
        _revoke_remember_me_token(user_id, new_raw_token)
    return None


def _safe_non_negative_decimal(value):
    try:
        return max(Decimal(str(value or 0)), Decimal("0"))
    except (InvalidOperation, ValueError, TypeError):
        return Decimal("0")


def calculate_budget_utilization(budget_amount, spent_amount):
    budget = _safe_non_negative_decimal(budget_amount)
    spent = _safe_non_negative_decimal(spent_amount)
    if budget <= 0:
        return None
    return spent / budget * Decimal("100")


def classify_budget_status(utilization_percentage):
    if utilization_percentage is None:
        return "Unavailable"
    if utilization_percentage < BUDGET_HEALTHY_LIMIT:
        return "Healthy"
    if utilization_percentage <= BUDGET_APPROACHING_LIMIT:
        return "Approaching Limit"
    return "Over Budget"


def build_budget_spending_analysis(budgets, stats):
    analysis_rows = []
    for budget in budgets:
        budget_amount = _safe_non_negative_decimal(budget.get("budget_amount"))
        spent_amount = _safe_non_negative_decimal(budget.get("spent_amount"))
        remaining_value = budget.get("remaining_amount")
        if remaining_value in (None, ""):
            remaining_amount = budget_amount - spent_amount
        else:
            try:
                remaining_amount = Decimal(str(remaining_value))
            except (InvalidOperation, ValueError, TypeError):
                remaining_amount = budget_amount - spent_amount

        utilization = calculate_budget_utilization(budget_amount, spent_amount)
        analysis_rows.append(
            {
                "budget_name": budget.get("budget_name") or "Budget",
                "category": budget.get("category") or "Uncategorized",
                "budget_amount": budget_amount,
                "spent_amount": spent_amount,
                "remaining_amount": remaining_amount,
                "utilization_percentage": utilization,
                "progress_percentage": min(utilization or Decimal("0"), Decimal("100")),
                "status": classify_budget_status(utilization),
            }
        )

    total_budget = _safe_non_negative_decimal(stats.get("total_allocated"))
    total_spent = _safe_non_negative_decimal(stats.get("total_spent"))
    total_remaining = stats.get("total_remaining")
    if total_remaining in (None, ""):
        total_remaining = total_budget - total_spent
    else:
        try:
            total_remaining = Decimal(str(total_remaining))
        except (InvalidOperation, ValueError, TypeError):
            total_remaining = total_budget - total_spent

    overall_utilization = calculate_budget_utilization(total_budget, total_spent)
    return analysis_rows, {
        "total_budget": total_budget,
        "total_spent": total_spent,
        "total_remaining": total_remaining,
        "utilization_percentage": overall_utilization,
        "status": classify_budget_status(overall_utilization),
    }


def _finite_decimal(value):
    if value in (None, ""):
        return None
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return None
    return number if number.is_finite() else None


def build_spending_recommendations(
    budget_spending_analysis,
    expense_breakdown,
    monthly_expenses,
    expense_stats,
):
    recommendations = []

    for budget in budget_spending_analysis or []:
        budget_name = budget.get("budget_name") or "Budget"
        utilization = _finite_decimal(budget.get("utilization_percentage"))
        status = budget.get("status")

        if status == "Over Budget" or (
            utilization is not None and utilization > OVER_BUDGET_THRESHOLD
        ):
            recommendations.append(
                {
                    "priority": "High",
                    "style": "danger",
                    "icon": "bi bi-exclamation-octagon-fill",
                    "title": "Budget Over Limit",
                    "title_key": "Budget Over Limit",
                    "message_key": "Your {budget_name} budget is over its stored limit. Review spending in this budget.",
                    "message_value": budget_name,
                    "message": (
                        f"Your {budget_name} budget is over its stored limit. "
                        "Review spending in this budget."
                    ),
                }
            )
        elif (
            utilization is not None
            and APPROACHING_BUDGET_THRESHOLD <= utilization <= OVER_BUDGET_THRESHOLD
        ):
            recommendations.append(
                {
                    "priority": "Medium",
                    "style": "warning",
                    "icon": "bi bi-exclamation-triangle-fill",
                    "title": "Budget Near Limit",
                    "title_key": "Budget Near Limit",
                    "message_key": "Your {budget_name} budget is approaching its stored limit. Review remaining spending.",
                    "message_value": budget_name,
                    "message": (
                        f"Your {budget_name} budget is approaching its stored limit. "
                        "Review remaining spending."
                    ),
                }
            )

    category_totals = []
    for item in expense_breakdown or []:
        amount = _finite_decimal(item.get("amount"))
        if amount is not None and amount > 0:
            category_totals.append((item.get("category") or "Other", amount))

    total_category_expenses = sum(
        (amount for _, amount in category_totals), Decimal("0")
    )
    if category_totals and total_category_expenses > 0:
        largest_category, largest_category_amount = max(
            category_totals, key=lambda item: item[1]
        )
        if (
            largest_category_amount / total_category_expenses
            >= CATEGORY_CONCENTRATION_THRESHOLD
        ):
            recommendations.append(
                {
                    "priority": "Medium",
                    "style": "info",
                    "icon": "bi bi-pie-chart-fill",
                    "title": "Expense Concentration",
                    "title_key": "Expense Concentration",
                    "message_key": "A large share of your recorded expenses is concentrated in {category}. Review this category for possible reductions.",
                    "message_value": largest_category,
                    "message": (
                        "A large share of your recorded expenses is concentrated in "
                        f"{largest_category}. Review this category for possible reductions."
                    ),
                }
            )

    if len(monthly_expenses or []) >= 2:
        previous_amount = _finite_decimal(monthly_expenses[-2].get("amount"))
        latest_amount = _finite_decimal(monthly_expenses[-1].get("amount"))
        if (
            previous_amount is not None
            and latest_amount is not None
            and latest_amount > previous_amount
        ):
            recommendations.append(
                {
                    "priority": "Medium",
                    "style": "warning",
                    "icon": "bi bi-graph-up-arrow",
                    "title": "Spending Increased",
                    "title_key": "Spending Increased",
                    "message_key": "Recorded spending increased compared with the previous recorded month. Review recent expenses.",
                    "message_value": "",
                    "message": (
                        "Recorded spending increased compared with the previous "
                        "recorded month. Review recent expenses."
                    ),
                }
            )

    total_expenses = _finite_decimal((expense_stats or {}).get("total_expenses"))
    average_expense = _finite_decimal((expense_stats or {}).get("average_expense"))
    largest_expense = _finite_decimal((expense_stats or {}).get("largest_expense"))
    if (
        total_expenses is not None
        and total_expenses > 0
        and average_expense is not None
        and average_expense > 0
        and largest_expense is not None
        and largest_expense >= LARGE_EXPENSE_MULTIPLIER * average_expense
    ):
        recommendations.append(
            {
                "priority": "Low",
                "style": "info",
                "icon": "bi bi-receipt-cutoff",
                "title": "Large Expense",
                "title_key": "Large Expense",
                "message_key": "One recorded expense is substantially larger than your average expense. Review that transaction.",
                "message_value": "",
                "message": (
                    "One recorded expense is substantially larger than your average "
                    "expense. Review that transaction."
                ),
            }
        )

    priority_order = {"High": 0, "Medium": 1, "Low": 2}
    recommendations.sort(key=lambda item: priority_order[item["priority"]])
    return recommendations


def investment_csrf_token():
    token = session.get("investment_csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["investment_csrf_token"] = token
    return token


def investment_csrf_valid():
    expected = session.get("investment_csrf_token")
    supplied = request.form.get("_investment_csrf_token", "")
    return bool(expected and supplied and secrets.compare_digest(expected, supplied))


def goal_csrf_token():
    token = session.get("goal_csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["goal_csrf_token"] = token
    return token


def goal_csrf_valid():
    expected = session.get("goal_csrf_token")
    supplied = request.form.get("_goal_csrf_token", "")
    return bool(expected and supplied and secrets.compare_digest(expected, supplied))


def budget_csrf_token():
    token = session.get("budget_csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["budget_csrf_token"] = token
    return token


def budget_csrf_valid():
    expected = session.get("budget_csrf_token")
    supplied = request.form.get("_budget_csrf_token", "")
    return bool(expected and supplied and secrets.compare_digest(expected, supplied))


def expense_csrf_token():
    token = session.get("expense_csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["expense_csrf_token"] = token
    return token


def expense_csrf_valid():
    expected = session.get("expense_csrf_token")
    supplied = request.form.get("_expense_csrf_token", "")
    return bool(expected and supplied and secrets.compare_digest(expected, supplied))


def preferences_csrf_token():
    token = session.get("preferences_csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["preferences_csrf_token"] = token
    return token


def preferences_csrf_valid():
    expected = session.get("preferences_csrf_token")
    supplied = request.form.get("_preferences_csrf_token", "")
    return bool(expected and supplied and secrets.compare_digest(expected, supplied))


def notifications_csrf_token():
    token = session.get("notifications_csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["notifications_csrf_token"] = token
    return token


def notifications_csrf_valid():
    expected = session.get("notifications_csrf_token")
    supplied = request.form.get("_notifications_csrf_token", "")
    return bool(expected and supplied and secrets.compare_digest(expected, supplied))


def notification_redirect_target():
    target = request.form.get("next", "")
    return target if target.startswith("/") and not target.startswith("//") else url_for("notifications")


@app.context_processor
def inject_notification_header_data():
    if not session.get("uid"):
        return {
            "notification_preview": [],
            "notification_unread_count": 0,
            "notification_csrf_token": "",
            "notification_theme_preferences": {},
            "notification_preferences_csrf_token": "",
            "current_currency": "USD",
            "current_currency_symbol": currency_symbol("USD"),
            "current_language": "en",
            "translation_catalog": TRANSLATIONS,
            "translate": translate,
            "t": lambda key: translate(key, "en"),
            "format_currency": format_currency,
        }
    try:
        user_id = current_user_id()
        preferences = get_user_preferences(user_id) or PREFERENCE_DEFAULTS
        return {
            "notification_preview": get_notifications(user_id, "all", 5),
            "notification_unread_count": get_unread_notification_count(user_id),
            "notification_csrf_token": notifications_csrf_token(),
            "notification_theme_preferences": preferences,
            "notification_preferences_csrf_token": preferences_csrf_token(),
            "current_currency": preferences.get("currency", "USD"),
            "current_currency_symbol": currency_symbol(preferences.get("currency")),
            "current_language": preferences.get("language", "en"),
            "translation_catalog": TRANSLATIONS,
            "translate": translate,
            "t": lambda key: translate(key, preferences.get("language", "en")),
            "format_currency": format_currency,
        }
    except Exception:
        app.logger.exception("Unable to load notification header data")
        return {
            "notification_preview": [],
            "notification_unread_count": 0,
            "notification_csrf_token": "",
            "notification_theme_preferences": {},
            "notification_preferences_csrf_token": "",
            "current_currency": "USD",
            "current_currency_symbol": currency_symbol("USD"),
            "current_language": "en",
            "translation_catalog": TRANSLATIONS,
            "translate": translate,
            "t": lambda key: translate(key, "en"),
            "format_currency": format_currency,
        }


def auth_csrf_token():
    token = session.get("auth_csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["auth_csrf_token"] = token
    return token


def auth_csrf_valid():
    expected = session.get("auth_csrf_token")
    supplied = request.form.get("_auth_csrf_token", "")
    return bool(expected and supplied and secrets.compare_digest(expected, supplied))


def _password_reset_challenge_id(session_key="password_reset_challenge_id"):
    challenge_id = session.get(session_key)
    return challenge_id if type(challenge_id) is int and challenge_id > 0 else None


def _otp_resend_view_data(resend_status):
    """Convert persisted resend state into safe data attributes for the OTP UI."""
    if not resend_status:
        return {
            "resend_remaining_seconds": 0,
            "resend_eligible": False,
            "resend_limit_reached": False,
        }
    try:
        resend_count = max(0, int(resend_status.get("resend_count", 0)))
        remaining_seconds = max(
            0, int(resend_status.get("remaining_seconds", 0) or 0)
        )
    except (TypeError, ValueError):
        resend_count = 0
        remaining_seconds = 0
    return {
        "resend_remaining_seconds": remaining_seconds,
        "resend_eligible": resend_count < 5
        and (remaining_seconds > 0 or bool(resend_status.get("can_resend"))),
        "resend_limit_reached": resend_count >= 5,
    }


def _password_reset_otp_hash(email):
    otp = f"{secrets.randbelow(1000000):06d}"
    try:
        otp_hash = registration_otp_digest(otp)
        email_service.send_password_reset_otp(email, otp)
    except Exception:
        app.logger.exception("Password reset verification email delivery failed")
        return None
    return otp_hash


def _password_reset_expiry():
    return datetime.now(timezone.utc) + timedelta(
        seconds=PASSWORD_RESET_OTP_TIMEOUT_SECONDS
    )


def _prepare_password_reset_for_user(user):
    challenge = get_password_reset_challenge(user["id"])
    if challenge:
        if challenge.get("verified_at") is not None:
            authorization = get_password_reset_authorization(challenge["id"])
            if authorization:
                session.pop("password_reset_challenge_id", None)
                session["password_reset_authorized_challenge_id"] = challenge["id"]
                return
            challenge = None

        if challenge:
            resend_status = get_password_reset_resend_status(challenge["id"])
            if resend_status and resend_status.get("otp_expired"):
                challenge = None
            elif not resend_status or not resend_status.get("can_resend"):
                session["password_reset_challenge_id"] = challenge["id"]
                return
            else:
                session["password_reset_challenge_id"] = challenge["id"]
                otp_hash = _password_reset_otp_hash(user["email"])
                if otp_hash:
                    updated = replace_password_reset_otp(
                        challenge["id"], otp_hash, _password_reset_expiry()
                    )
                    if updated:
                        session["password_reset_challenge_id"] = updated["id"]
                return

    otp_hash = _password_reset_otp_hash(user["email"])
    if not otp_hash:
        return
    challenge = create_password_reset_challenge(
        user["id"], user["email"], otp_hash, _password_reset_expiry()
    )
    if not challenge:
        challenge = get_password_reset_challenge(user["id"])
    if challenge:
        session["password_reset_challenge_id"] = challenge["id"]


app.jinja_env.globals["auth_csrf_token"] = auth_csrf_token


def security_sessions_csrf_token():
    token = session.get("security_sessions_csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["security_sessions_csrf_token"] = token
    return token


def security_sessions_csrf_valid():
    expected = session.get("security_sessions_csrf_token")
    supplied = request.form.get("_security_sessions_csrf_token", "")
    return bool(expected and supplied and secrets.compare_digest(expected, supplied))


def two_factor_login_csrf_token():
    token = session.get("two_factor_login_csrf_token")
    if not token:
        token = secrets.token_urlsafe(32)
        session["two_factor_login_csrf_token"] = token
    return token


def two_factor_login_csrf_valid():
    expected = session.get("two_factor_login_csrf_token")
    supplied = request.form.get("_two_factor_login_csrf_token", "")
    return bool(expected and supplied and secrets.compare_digest(expected, supplied))


@app.route("/")
def home():
    if session.get("uid"):
        return redirect(url_for("dashboard"))
    return render_template(
        "landing.html", current_year=datetime.now(timezone.utc).year
    )


@app.get("/health")
def health():
    return jsonify(status="ok")


@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("uid"):
        if login_required_redirect() is None:
            return redirect(url_for("dashboard"))

    if request.method == "POST":
        if not auth_csrf_valid():
            return render_template(
                "login.html", error="The form security token is missing or invalid."
            ), 400
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        remember_me = _remember_me_requested()

        if _login_rate_limited():
            return render_template(
                "login.html",
                error="Too many failed login attempts. Please try again later.",
            ), 429

        success, user = login_user(email, password)
        if success:
            _clear_login_failures()
            session.clear()
            if get_totp_status(user["id"])["is_enabled"]:
                session["pending_2fa_user_id"] = user["id"]
                session["pending_2fa_remember_me"] = remember_me
                return redirect(url_for("login_two_factor"))

            if not establish_authenticated_session(user):
                return render_template(
                    "login.html",
                    error="Unable to start a secure session. Please try again.",
                ), 503
            response = redirect(url_for("dashboard"))
            if remember_me:
                raw_token = _issue_remember_me_token(user["id"])
                if raw_token:
                    _set_remember_me_cookie(response, raw_token)
            return response

        _record_login_failure()
        return render_template("login.html", error="Invalid email or password.")

    success = (
        "Password reset successful. Please sign in with your new password."
        if request.args.get("reset") == "success"
        else None
    )
    return render_template("login.html", error=None, success=success)


@app.route("/login/2fa", methods=["GET", "POST"])
def login_two_factor():
    pending_user_id = session.get("pending_2fa_user_id")
    if not pending_user_id:
        return redirect(url_for("login"))

    if request.method == "GET":
        return render_template(
            "two_factor_login.html",
            two_factor_login_csrf_token=two_factor_login_csrf_token(),
            error=None,
        )

    if not two_factor_login_csrf_valid():
        return render_template(
            "two_factor_login.html",
            two_factor_login_csrf_token=two_factor_login_csrf_token(),
            error="The form security token is missing or invalid.",
        ), 400
    if _login_rate_limited():
        return render_template(
            "two_factor_login.html",
            two_factor_login_csrf_token=two_factor_login_csrf_token(),
            error="Too many failed verification attempts. Please try again later.",
        ), 429

    credential = get_totp_credential(pending_user_id)
    if not credential or credential["enabled_at"] is None:
        session.clear()
        return redirect(url_for("login"))

    try:
        is_valid = pyotp.TOTP(
            decrypt_totp_secret(credential["secret_encrypted"])
        ).verify(request.form.get("totp_code", ""), valid_window=0)
    except (InvalidToken, ValueError):
        is_valid = False

    if not is_valid:
        _record_login_failure()
        return render_template(
            "two_factor_login.html",
            two_factor_login_csrf_token=two_factor_login_csrf_token(),
            error="Invalid verification code.",
        )

    remember_me = bool(session.pop("pending_2fa_remember_me", False))
    remembered_session = bool(session.pop("pending_2fa_remembered_session", False))
    user = get_user_by_id(pending_user_id)
    if not user or not establish_authenticated_session(user):
        session.clear()
        return render_template(
            "login.html",
            error="Unable to start a secure session. Please try again.",
        ), 503

    _clear_login_failures()
    response = redirect(url_for("dashboard"))
    if remember_me and not remembered_session:
        raw_token = _issue_remember_me_token(user["id"])
        if raw_token:
            _set_remember_me_cookie(response, raw_token)
    return response


@app.route("/forgot-password", methods=["GET", "POST"])
def forgot_password():
    if request.method == "GET":
        return render_template("forgot_password.html", error=None, success=None)

    if not auth_csrf_valid():
        return render_template(
            "forgot_password.html",
            error="The form security token is missing or invalid.",
            success=None,
        ), 400

    email = request.form.get("email", "").strip().lower()
    if not email or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        return render_template(
            "forgot_password.html",
            error="Enter a valid email address.",
            success=None,
        ), 400

    session.pop("password_reset_challenge_id", None)
    session.pop("password_reset_authorized_challenge_id", None)
    try:
        user = get_user_for_password_reset(email)
        if user:
            _prepare_password_reset_for_user(user)
    except Exception:
        app.logger.exception("Password reset request failed")

    return redirect(url_for("forgot_password_verify"))


def render_password_reset_otp(error=None, success=None, status=200):
    challenge_id = _password_reset_challenge_id()
    challenge = (
        get_password_reset_challenge_by_id(challenge_id) if challenge_id else None
    )
    resend_status = (
        get_password_reset_resend_status(challenge_id) if challenge else None
    )
    if not resend_status:
        resend_status = {"resend_count": 0, "remaining_seconds": 60}
    return render_template(
        "password_reset_otp.html",
        error=error,
        success=success,
        auth_csrf_token=auth_csrf_token(),
        **_otp_resend_view_data(resend_status),
    ), status


@app.route("/forgot-password/verify", methods=["GET", "POST"])
def forgot_password_verify():
    authorized_id = _password_reset_challenge_id(
        "password_reset_authorized_challenge_id"
    )
    if request.method == "GET":
        if authorized_id and get_password_reset_authorization(authorized_id):
            return redirect(url_for("password_reset"))
        return render_password_reset_otp(
            success=PASSWORD_RESET_GENERIC_MESSAGE,
        )

    if not auth_csrf_valid():
        return render_password_reset_otp(
            error="The form security token is missing or invalid.", status=400
        )

    challenge_id = _password_reset_challenge_id()
    otp = request.form.get("otp", "").strip()
    if not challenge_id or not otp.isdigit() or len(otp) != 6:
        return render_password_reset_otp(
            error="Enter the six-digit verification code.", status=400
        )

    reset_token = secrets.token_urlsafe(32)
    try:
        verified, result = verify_password_reset_otp(
            challenge_id,
            registration_otp_digest(otp),
            hashlib.sha256(reset_token.encode("utf-8")).hexdigest(),
            datetime.now(timezone.utc)
            + timedelta(seconds=PASSWORD_RESET_AUTH_TIMEOUT_SECONDS),
        )
    except Exception:
        app.logger.exception("Password reset verification failed")
        return render_password_reset_otp(
            error="Unable to verify the code. Please try again.", status=503
        )

    if not verified:
        messages = {
            "expired": "This verification code has expired. Please request a new code.",
            "attempts": "Too many verification attempts. Please request a new code.",
            "consumed": "This verification code is no longer valid. Please request a new code.",
        }
        return render_password_reset_otp(
            error=messages.get(result, PASSWORD_RESET_INVALID_MESSAGE), status=400
        )

    session.pop("password_reset_challenge_id", None)
    session["password_reset_authorized_challenge_id"] = challenge_id
    return redirect(url_for("password_reset"))


@app.post("/forgot-password/resend")
def resend_password_reset_otp():
    if not auth_csrf_valid():
        return render_password_reset_otp(
            error="The form security token is missing or invalid.", status=400
        )

    challenge_id = _password_reset_challenge_id()
    challenge = (
        get_password_reset_challenge_by_id(challenge_id) if challenge_id else None
    )
    status = (
        get_password_reset_resend_status(challenge_id) if challenge else None
    )
    if not status or not status.get("can_resend"):
        error = (
            "You have reached the resend limit. Please request a new password reset."
            if status and status.get("resend_count", 0) >= 5
            else "Please wait before requesting another code."
        )
        return render_password_reset_otp(error=error, status=429)

    otp_hash = _password_reset_otp_hash(challenge["email"])
    if not otp_hash:
        return render_password_reset_otp(
            error="Unable to send the verification email. Please try again later.",
            status=503,
        )

    updated = replace_password_reset_otp(
        challenge_id, otp_hash, _password_reset_expiry()
    )
    if not updated:
        return render_password_reset_otp(
            error="Please wait before requesting another code.", status=429
        )
    return render_password_reset_otp(
        success="A new verification code has been sent."
    )


@app.route("/forgot-password/reset", methods=["GET", "POST"])
def password_reset():
    challenge_id = _password_reset_challenge_id(
        "password_reset_authorized_challenge_id"
    )
    authorization = (
        get_password_reset_authorization(challenge_id) if challenge_id else None
    )
    if not authorization:
        session.pop("password_reset_authorized_challenge_id", None)
        return redirect(url_for("forgot_password_verify"))

    if request.method == "GET":
        return render_template("password_reset.html", error=None)

    if not auth_csrf_valid():
        return render_template(
            "password_reset.html",
            error="The form security token is missing or invalid.",
        ), 400

    new_password = request.form.get("new_password", "")
    confirm_password = request.form.get("confirm_password", "")
    if len(new_password) < 8:
        return render_template(
            "password_reset.html",
            error="Password must contain at least 8 characters.",
        ), 400
    if new_password != confirm_password:
        return render_template(
            "password_reset.html",
            error="Passwords do not match.",
        ), 400

    try:
        changed = reset_password_with_authorization(
            challenge_id, generate_password_hash(new_password)
        )
    except Exception:
        app.logger.exception("Password reset failed")
        changed = False
    if not changed:
        session.pop("password_reset_authorized_challenge_id", None)
        return render_template(
            "password_reset.html",
            error="Unable to reset password. Please start again.",
        ), 400

    session.clear()
    return redirect(url_for("login", reset="success"))


@app.route("/register", methods=["GET", "POST"])
def register():
    if session.get("uid"):
        if login_required_redirect() is None:
            return redirect(url_for("dashboard"))

    if request.method == "GET":
        return redirect(url_for("login"))

    if not auth_csrf_valid():
        return render_template(
            "login.html", error="The form security token is missing or invalid."
        ), 400

    full_name = request.form.get("username", "").strip()
    email = request.form.get("email", "").strip().lower()
    mobile_number = request.form.get("mobile_number", "").strip()
    password = request.form.get("password", "")
    confirm_password = request.form.get("confirm_password", "")

    if not password:
        return render_template("login.html", error="All fields are required.")
    if len(password) < 8:
        return render_template(
            "login.html", error="Password must contain at least 8 characters."
        )
    if not full_name or not email or not mobile_number or not confirm_password:
        return render_template("login.html", error="All fields are required.")
    if len(full_name) > 150 or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        return render_template("login.html", error="Enter valid registration details.")
    if not re.fullmatch(r"[0-9+() -]{7,30}", mobile_number):
        return render_template("login.html", error="Enter a valid mobile number.")
    if password != confirm_password:
        return render_template("login.html", error="Passwords do not match.")

    otp = f"{secrets.randbelow(1000000):06d}"
    try:
        otp_hash = registration_otp_digest(otp)
    except RuntimeError:
        return render_template(
            "login.html",
            error="Unable to begin registration verification.",
        ), 503
    pending = create_pending_registration(
        full_name,
        email,
        mobile_number,
        generate_password_hash(password),
        otp_hash,
        datetime.now(timezone.utc) + timedelta(minutes=10),
    )
    if not pending:
        return render_template("login.html", error="Unable to begin registration verification.")
    try:
        email_service.send_registration_otp(email, otp)
    except Exception:
        app.logger.exception("Registration verification email delivery failed")
        return render_template(
            "login.html",
            error="Unable to send the verification email. Please try again later.",
        ), 503

    session["pending_registration_id"] = pending["id"]
    return render_registration_otp(success="A verification code was sent to your email.")


def render_registration_otp(error=None, success=None, status=200):
    pending_id = session.get("pending_registration_id")
    pending = get_pending_registration(pending_id) if pending_id else None
    if not pending:
        return redirect(url_for("login"))
    resend_status = get_registration_resend_status(pending_id)
    return render_template(
        "registration_otp.html",
        email=pending["email"],
        auth_csrf_token=auth_csrf_token(),
        error=error,
        success=success,
        **_otp_resend_view_data(resend_status),
    ), status


@app.route("/verify-registration-otp", methods=["GET", "POST"])
def verify_registration_otp():
    if request.method == "GET":
        return render_registration_otp()
    if not auth_csrf_valid():
        return render_registration_otp(
            "The form security token is missing or invalid.", status=400
        )

    pending_id = session.get("pending_registration_id")
    otp = request.form.get("otp", "").strip()
    if not isinstance(pending_id, int) or not otp.isdigit() or len(otp) != 6:
        return render_registration_otp(
            "Enter the six-digit verification code.", status=400
        )

    success, result = complete_pending_registration(
        pending_id,
        registration_otp_digest(otp),
    )
    if not success:
        messages = {
            "expired": "This verification code has expired. Please request a new code.",
            "attempts": "Too many verification attempts. Please request a new code.",
        }
        return render_registration_otp(messages.get(result, "The verification code is invalid."))

    session.pop("pending_registration_id", None)
    return render_template("registration_success.html")


@app.post("/resend-registration-otp")
def resend_registration_otp():
    if not auth_csrf_valid():
        return render_registration_otp(
            "The form security token is missing or invalid.", status=400
        )
    pending_id = session.get("pending_registration_id")
    pending = get_pending_registration(pending_id) if isinstance(pending_id, int) else None
    if not pending:
        return redirect(url_for("login"))

    resend_status = get_registration_resend_status(pending_id)
    if not resend_status or not resend_status.get("can_resend"):
        message = (
            "You have reached the resend limit for this registration."
            if resend_status and resend_status.get("resend_count", 0) >= 5
            else "Please wait before requesting another code."
        )
        return render_registration_otp(message, status=429)

    otp = f"{secrets.randbelow(1000000):06d}"
    try:
        otp_hash = registration_otp_digest(otp)
    except RuntimeError:
        return render_registration_otp(
            "Unable to send the verification email. Please try again later.", status=503
        )
    try:
        email_service.send_registration_otp(pending["email"], otp)
    except Exception:
        app.logger.exception("Registration verification email delivery failed")
        return render_registration_otp(
            "Unable to send the verification email. Please try again later.", status=503
        )
    updated = replace_pending_registration_otp(
        pending_id,
        otp_hash,
        datetime.now(timezone.utc) + timedelta(minutes=10),
    )
    if not updated:
        return render_registration_otp(
            "Unable to update the verification challenge. Please try again.", status=503
        )
    return render_registration_otp(success="A new verification code was sent.")


@app.post("/logout")
def logout():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response
    if not auth_csrf_valid():
        return "The form security token is missing or invalid.", 400
    user_id = session.get("uid")
    session_token_hash = current_session_token_hash()
    remember_me_token = request.cookies.get(REMEMBER_ME_COOKIE_NAME)
    rotated_remember_me_token = getattr(g, "remember_me_cookie_token", None)
    if user_id and session_token_hash:
        try:
            revoke_current_user_session(user_id, session_token_hash)
        except Exception:
            app.logger.exception("Authenticated session revocation failed during logout")
    _revoke_remember_me_token(user_id, remember_me_token)
    _revoke_remember_me_token(user_id, rotated_remember_me_token)
    session.clear()
    return _clear_remember_me_cookie(redirect(url_for("login")))


@app.get("/profile")
def profile():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    user = get_user_by_id(current_user_id())
    if not user:
        flash("Profile not found.", "danger")
        return redirect(url_for("dashboard"))

    return render_template("profile/dashboard.html", user=user)


def security_page_data(user_id):
    return (
        list_active_user_sessions(user_id, current_session_token_hash() or ""),
        get_totp_status(user_id),
    )


def render_security_sessions(
    active_sessions, totp_status, status=200, setup_secret=None
):
    return render_template(
        "profile/security.html",
        active_sessions=active_sessions,
        totp_status=totp_status,
        setup_secret=setup_secret,
        security_sessions_csrf_token=security_sessions_csrf_token(),
    ), status


@app.get("/profile/security")
def profile_security():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    active_sessions, totp_status = security_page_data(current_user_id())
    return render_security_sessions(active_sessions, totp_status)


@app.post("/profile/security/password")
def change_password():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    user_id = current_user_id()
    active_sessions, totp_status = security_page_data(user_id)
    if not security_sessions_csrf_valid():
        flash("The form security token is missing or invalid.", "danger")
        return render_security_sessions(active_sessions, totp_status, 400)

    current_password = request.form.get("current_password", "")
    new_password = request.form.get("new_password", "")
    confirm_password = request.form.get("confirm_password", "")
    if not current_password or not new_password or not confirm_password:
        flash("All password fields are required.", "danger")
        return render_security_sessions(active_sessions, totp_status, 400)
    if not verify_user_password(user_id, current_password):
        flash("Current password is incorrect.", "danger")
        return render_security_sessions(active_sessions, totp_status, 400)
    if new_password != confirm_password:
        flash("New password and confirmation do not match.", "danger")
        return render_security_sessions(active_sessions, totp_status, 400)
    if verify_user_password(user_id, new_password):
        flash("New password must be different from the current password.", "danger")
        return render_security_sessions(active_sessions, totp_status, 400)
    if len(new_password) < 8:
        flash("Password must contain at least 8 characters.", "danger")
        return render_security_sessions(active_sessions, totp_status, 400)

    try:
        changed = update_user_password_and_revoke_other_sessions(
            user_id, new_password, current_session_token_hash() or ""
        )
    except Exception:
        app.logger.exception("Password change failed")
        flash("Unable to change password. Please try again.", "danger")
        return render_security_sessions(active_sessions, totp_status, 503)

    if not changed:
        flash("Unable to change password. Please try again.", "danger")
        return render_security_sessions(active_sessions, totp_status, 503)

    _forget_remember_me_cookie()
    flash("Password changed successfully. Other active devices were signed out.", "success")
    return redirect(url_for("profile_security"))


@app.post("/profile/security/totp/setup")
def setup_totp():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    user_id = current_user_id()
    active_sessions, totp_status = security_page_data(user_id)
    if not security_sessions_csrf_valid():
        flash("The form security token is missing or invalid.", "danger")
        return render_security_sessions(active_sessions, totp_status, 400)
    if totp_status["is_enabled"]:
        flash("Two-factor authentication is already enabled.", "danger")
        return render_security_sessions(active_sessions, totp_status, 400)

    secret = pyotp.random_base32()
    try:
        save_pending_totp_secret(user_id, encrypt_totp_secret(secret))
    except Exception:
        app.logger.exception("Two-factor setup could not be saved")
        flash("Unable to start two-factor setup. Please try again.", "danger")
        return render_security_sessions(active_sessions, totp_status, 503)

    active_sessions, totp_status = security_page_data(user_id)
    return render_security_sessions(
        active_sessions, totp_status, setup_secret=secret
    )


@app.post("/profile/security/totp/verify")
def verify_totp_setup():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    user_id = current_user_id()
    active_sessions, totp_status = security_page_data(user_id)
    if not security_sessions_csrf_valid():
        flash("The form security token is missing or invalid.", "danger")
        return render_security_sessions(active_sessions, totp_status, 400)

    credential = get_totp_credential(user_id)
    if not credential or credential["enabled_at"] is not None:
        flash("No pending two-factor setup was found.", "danger")
        return render_security_sessions(active_sessions, totp_status, 400)

    try:
        is_valid = pyotp.TOTP(
            decrypt_totp_secret(credential["secret_encrypted"])
        ).verify(request.form.get("totp_code", ""), valid_window=0)
    except (InvalidToken, ValueError):
        is_valid = False

    if not is_valid or not enable_totp_for_user(user_id):
        flash("Invalid verification code.", "danger")
        return render_security_sessions(active_sessions, totp_status, 400)

    flash("Two-factor authentication is enabled.", "success")
    return redirect(url_for("profile_security"))


@app.post("/profile/security/totp/disable")
def disable_totp():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    user_id = current_user_id()
    active_sessions, totp_status = security_page_data(user_id)
    if not security_sessions_csrf_valid():
        flash("The form security token is missing or invalid.", "danger")
        return render_security_sessions(active_sessions, totp_status, 400)
    if not totp_status["is_enabled"]:
        flash("Two-factor authentication is not enabled.", "danger")
        return render_security_sessions(active_sessions, totp_status, 400)
    if not verify_user_password(user_id, request.form.get("current_password", "")):
        flash("Current password is incorrect.", "danger")
        return render_security_sessions(active_sessions, totp_status, 400)

    credential = get_totp_credential(user_id)
    try:
        is_valid = credential and pyotp.TOTP(
            decrypt_totp_secret(credential["secret_encrypted"])
        ).verify(request.form.get("totp_code", ""), valid_window=0)
    except (InvalidToken, ValueError):
        is_valid = False

    if not is_valid:
        flash("Invalid verification code.", "danger")
        return render_security_sessions(active_sessions, totp_status, 400)
    if not disable_totp_for_user(user_id):
        flash("Unable to disable two-factor authentication.", "danger")
        return render_security_sessions(active_sessions, totp_status, 503)

    flash("Two-factor authentication is disabled.", "success")
    return redirect(url_for("profile_security"))


@app.post("/profile/security/sessions/<int:session_id>/logout")
def logout_device_session(session_id):
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    user_id = current_user_id()
    session_token_hash = current_session_token_hash() or ""
    if not security_sessions_csrf_valid():
        flash("The form security token is missing or invalid.", "danger")
        active_sessions, totp_status = security_page_data(user_id)
        return render_security_sessions(active_sessions, totp_status, 400)

    if revoke_user_session(session_id, user_id, session_token_hash):
        flash("Device signed out successfully.", "success")
    else:
        flash("Active device not found.", "danger")
    return redirect(url_for("profile_security"))


@app.post("/profile/security/sessions/logout-all")
def logout_all_device_sessions():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    if not security_sessions_csrf_valid():
        flash("The form security token is missing or invalid.", "danger")
        active_sessions, totp_status = security_page_data(current_user_id())
        return render_security_sessions(active_sessions, totp_status, 400)

    user_id = current_user_id()
    revoke_all_user_sessions(user_id)
    try:
        revoke_all_remember_me_tokens(user_id)
    except Exception:
        app.logger.exception("Remember Me credential revocation failed during logout all")
    session.clear()
    _forget_remember_me_cookie()
    flash("You have been signed out from all devices.", "success")
    return _clear_remember_me_cookie(redirect(url_for("login")))


def preference_reference_options():
    currency_options = list_preference_currencies()
    language_options = list_preference_languages()
    if not any(option.get("code") == "mr" for option in language_options):
        language_options.append({"code": "mr", "display_name": "Marathi"})
    return currency_options, language_options


def validate_preferences_form(form, currency_options, language_options):
    preferences = {
        "theme": form.get("theme", "").strip().lower(),
        "currency": form.get("currency", "").strip().upper(),
        "language": form.get("language", "").strip().lower(),
    }
    errors = []

    if preferences["theme"] not in PREFERENCE_THEME_OPTIONS:
        errors.append("Select a supported theme.")
    currency_codes = {option["code"] for option in currency_options}
    language_codes = {option["code"] for option in language_options}
    if preferences["currency"] not in currency_codes:
        errors.append("Select a supported currency.")
    if preferences["language"] not in language_codes:
        errors.append("Select a supported language.")

    for field in PREFERENCE_BOOLEAN_FIELDS:
        preferences[field] = field in form

    return preferences, errors


def render_preferences(preferences, currency_options, language_options, status=200):
    sync_theme_session(preferences)
    return render_template(
        "profile/preferences.html",
        preferences=preferences or PREFERENCE_DEFAULTS,
        theme_options=PREFERENCE_THEME_DISPLAY_OPTIONS,
        currency_options=currency_options,
        language_options=language_options,
        preferences_csrf_token=preferences_csrf_token(),
    ), status


@app.route("/profile/preferences", methods=["GET", "POST"])
def profile_preferences():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    user_id = current_user_id()
    currency_options, language_options = preference_reference_options()

    if request.method == "GET":
        return render_preferences(
            ensure_user_preferences(user_id), currency_options, language_options
        )

    if not preferences_csrf_valid():
        flash("The form security token is missing or invalid.", "danger")
        return render_preferences(
            get_user_preferences(user_id), currency_options, language_options, 400
        )

    updated_preferences, errors = validate_preferences_form(
        request.form, currency_options, language_options
    )
    if errors:
        for error in errors:
            flash(error, "danger")
        return render_preferences(
            get_user_preferences(user_id), currency_options, language_options, 400
        )

    preferences = ensure_user_preferences(user_id)

    try:
        saved_preferences = update_user_preferences(user_id, updated_preferences)
    except Exception:
        app.logger.exception("Preference update failed")
        flash("Unable to save preferences. Please try again.", "danger")
        return render_preferences(preferences, currency_options, language_options, 503)

    if not saved_preferences:
        flash("Unable to save preferences. Please try again.", "danger")
        return render_preferences(preferences, currency_options, language_options, 503)

    flash("Preferences saved successfully.", "success")
    return redirect(url_for("profile_preferences"))


@app.route("/dashboard")
def dashboard():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    user_id = current_user_id()
    # Existing M1/M2 dashboard data
    stats = get_summary_stats(user_id)

    expense_stats = get_expense_summary(user_id)
    expense_transactions = get_transactions(user_id)
    monthly_expenses = get_monthly_expense_summary(user_id)
    budget_rows = filter_budgets(user_id)
    budget_spending_analysis, budget_spending_summary = build_budget_spending_analysis(
        budget_rows, stats
    )

    expense_total = float(expense_stats.get("total_spent") or 0)
    category_totals = {}
    for transaction in expense_transactions:
        category = transaction.get("category") or "Other"
        category_totals[category] = category_totals.get(category, 0) + float(
            transaction.get("amount") or 0
        )

    expense_colors = [
        "#2563EB",
        "#22C55E",
        "#F59E0B",
        "#EF4444",
        "#8B5CF6",
        "#06B6D4",
    ]
    expense_breakdown = [
        {
            "category": category,
            "amount": amount,
            "percentage": (amount / expense_total * 100) if expense_total else 0,
            "color": expense_colors[index % len(expense_colors)],
        }
        for index, (category, amount) in enumerate(
            sorted(category_totals.items(), key=lambda item: item[1], reverse=True)
        )
    ]
    spending_recommendations = build_spending_recommendations(
        budget_spending_analysis,
        expense_breakdown,
        monthly_expenses,
        expense_stats,
    )

    budget_progress = []
    for budget in budget_rows:
        budget_amount = float(budget.get("budget_amount") or 0)
        spent_amount = float(budget.get("spent_amount") or 0)
        budget_progress.append(
            {
                "budget_name": budget.get("budget_name") or "Budget",
                "budget_amount": budget_amount,
                "spent_amount": spent_amount,
                "remaining_amount": float(budget.get("remaining_amount") or 0),
                "percentage": min(
                    100,
                    (spent_amount / budget_amount * 100) if budget_amount else 0,
                ),
            }
        )
    # Existing M2 goal data
    goals = goal_service.list_for_user(user_id)

    # Existing M2 investment analytics
    _, investment_stats = investment_service.list_for_user(user_id)

    # Aggregate goal progress using existing goal values
    goal_total_current = sum(
        Decimal(str(goal.get("current_amount") or 0))
        for goal in goals
    )
    goal_total_target = sum(
        Decimal(str(goal.get("target_amount") or 0))
        for goal in goals
    )

    # Calculate M3 Financial Health Score
    financial_health = calculate_financial_health(
        budget_data={
            "total_allocated": stats.get("total_allocated"),
            "total_spent": stats.get("total_spent"),
        },
        spending_data={
            "total_spent": expense_stats.get("total_spent"),
            "average_expense": expense_stats.get("average_expense"),
            "largest_expense": expense_stats.get("largest_expense"),
            "expense_count": expense_stats.get("total_expenses"),
            "month_spent": expense_stats.get("month_spent"),
            "month_expenses": expense_stats.get("month_expenses"),
        },
        goal_data={
            "total_current": goal_total_current,
            "total_target": goal_total_target,
        },
        investment_data={
            "return_percentage": investment_stats.get("return_percentage"),
        },
    )
    return render_template(
        "dashboard.html",
        stats=stats,
        financial_health=financial_health,
        dashboard_month_spent=float(expense_stats.get("month_spent") or 0),
        dashboard_budget_left=float(stats.get("total_remaining") or 0),
        expense_breakdown=expense_breakdown,
        monthly_expenses=monthly_expenses,
        budget_spending_analysis=budget_spending_analysis,
        budget_spending_summary=budget_spending_summary,
        spending_recommendations=spending_recommendations,
        budget_progress=budget_progress,
        recent_transactions=expense_transactions[:5],
        current_username=session["username"],
    )


@app.get("/notifications")
def notifications():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    selected_filter = request.args.get("filter", "all").strip().lower()
    if selected_filter not in {"all", "unread", "alert", "milestone"}:
        selected_filter = "all"
    return render_template(
        "notifications/dashboard.html",
        notifications=get_notifications(current_user_id(), selected_filter),
        selected_filter=selected_filter,
        notifications_csrf_token=notifications_csrf_token(),
    )


@app.post("/notifications/<int:notification_id>/read")
def notification_read(notification_id):
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response
    if not notifications_csrf_valid():
        flash("The form security token is missing or invalid.", "danger")
        return redirect(url_for("notifications")), 400
    mark_notification_read(current_user_id(), notification_id)
    return redirect(notification_redirect_target())


@app.post("/notifications/read-all")
def notifications_read_all():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response
    if not notifications_csrf_valid():
        flash("The form security token is missing or invalid.", "danger")
        return redirect(url_for("notifications")), 400
    mark_all_notifications_read(current_user_id())
    return redirect(notification_redirect_target())


@app.get("/reports")
def reports_data():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    accepts = request.headers.get("Accept", "").lower()
    wants_json = (
        request.args.get("format") == "json"
        or not accepts
        or ("application/json" in accepts and "text/html" not in accepts)
    )

    try:
        report_data = build_reporting_data(
            current_user_id(),
            request.args.get("start_date"),
            request.args.get("end_date"),
            investment_service=investment_service,
            goal_service=goal_service,
        )
    except ReportValidationError as error:
        if wants_json:
            return jsonify({"error": str(error)}), 400
        return render_template(
            "reports/dashboard.html",
            report=None,
            error=str(error),
            start_date=request.args.get("start_date", ""),
            end_date=request.args.get("end_date", ""),
        ), 400

    if wants_json:
        return jsonify(report_data)
    return render_template(
        "reports/dashboard.html",
        report=report_data,
        start_date=request.args.get("start_date", ""),
        end_date=request.args.get("end_date", ""),
    )


def _pdf_decimal(value):
    if value in (None, ""):
        return None
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return None
    return number if number.is_finite() else None


def _pdf_amount(value, currency="INR"):
    number = _pdf_decimal(value)
    return "Unavailable" if number is None else format_currency(number, currency)


def _pdf_percent(value):
    number = _pdf_decimal(value)
    return "Unavailable" if number is None else f"{number:,.2f}%"


def _pdf_number(value):
    number = _pdf_decimal(value)
    if number is None:
        return "Unavailable"
    if number == number.to_integral_value():
        return str(int(number))
    return f"{number:,.2f}"


def _pdf_count(value):
    number = _pdf_decimal(value)
    return "0" if number is None else str(int(number))


def _pdf_date(value):
    if not value:
        return "Unavailable"
    try:
        return datetime.strptime(str(value), "%Y-%m-%d").strftime("%B %d, %Y").replace(
            " 0", " "
        )
    except (TypeError, ValueError):
        return str(value)


def _pdf_period(report_data):
    date_range = report_data.get("date_range") or {}
    start_date = date_range.get("start_date")
    end_date = date_range.get("end_date")
    if start_date and end_date:
        return f"{_pdf_date(start_date)} - {_pdf_date(end_date)}"
    return "All Available Expense History"


def _register_pdf_fonts():
    families = [
        (
            r"C:\Windows\Fonts\arial.ttf",
            r"C:\Windows\Fonts\arialbd.ttf",
            r"C:\Windows\Fonts\ariali.ttf",
        ),
        (
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf",
        ),
        (
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Italic.ttf",
        ),
    ]
    for regular_path, bold_path, italic_path in families:
        if not os.path.exists(regular_path):
            continue
        try:
            pdfmetrics.registerFont(TTFont("FinSightPDFRegular", regular_path))
            bold_name = "FinSightPDFBold"
            italic_name = "FinSightPDFItalic"
            if os.path.exists(bold_path):
                pdfmetrics.registerFont(TTFont(bold_name, bold_path))
            else:
                bold_name = "FinSightPDFRegular"
            if os.path.exists(italic_path):
                pdfmetrics.registerFont(TTFont(italic_name, italic_path))
            else:
                italic_name = "FinSightPDFRegular"
            return "FinSightPDFRegular", bold_name, italic_name
        except (OSError, RuntimeError, ValueError):
            continue
    return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique"


PDF_FONT, PDF_FONT_BOLD, PDF_FONT_ITALIC = _register_pdf_fonts()


def _pdf_styles():
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "PDFTitle",
            parent=styles["Title"],
            fontName=PDF_FONT_BOLD,
            fontSize=24,
            leading=28,
            textColor=colors.HexColor("#091A2B"),
            alignment=TA_LEFT,
            spaceAfter=2,
        ),
        "subtitle": ParagraphStyle(
            "PDFSubtitle",
            parent=styles["Normal"],
            fontName=PDF_FONT,
            fontSize=13,
            leading=16,
            textColor=colors.HexColor("#138A70"),
            spaceAfter=12,
        ),
        "section": ParagraphStyle(
            "PDFSection",
            parent=styles["Heading2"],
            fontName=PDF_FONT_BOLD,
            fontSize=14,
            leading=18,
            textColor=colors.HexColor("#138A70"),
            spaceBefore=14,
            spaceAfter=7,
        ),
        "body": ParagraphStyle(
            "PDFBody",
            parent=styles["BodyText"],
            fontName=PDF_FONT,
            fontSize=9,
            leading=13,
            textColor=colors.HexColor("#122333"),
            spaceAfter=5,
        ),
        "muted": ParagraphStyle(
            "PDFMuted",
            parent=styles["BodyText"],
            fontName=PDF_FONT_ITALIC,
            fontSize=8.5,
            leading=12,
            textColor=colors.HexColor("#71808D"),
            spaceAfter=5,
        ),
        "table_header": ParagraphStyle(
            "PDFTableHeader",
            parent=styles["BodyText"],
            fontName=PDF_FONT_BOLD,
            fontSize=8,
            leading=10,
            textColor=colors.white,
        ),
        "table_body": ParagraphStyle(
            "PDFTableBody",
            parent=styles["BodyText"],
            fontName=PDF_FONT,
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#122333"),
        ),
        "summary_label": ParagraphStyle(
            "PDFSummaryLabel",
            parent=styles["BodyText"],
            fontName=PDF_FONT_BOLD,
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#71808D"),
            alignment=TA_CENTER,
        ),
        "summary_value": ParagraphStyle(
            "PDFSummaryValue",
            parent=styles["BodyText"],
            fontName=PDF_FONT_BOLD,
            fontSize=10,
            leading=13,
            textColor=colors.HexColor("#091A2B"),
            alignment=TA_CENTER,
        ),
    }


def _pdf_cell(value, style):
    text = "" if value is None else str(value)
    return Paragraph(escape(text), style)


def _pdf_table(rows, widths, styles):
    formatted_rows = []
    for index, row in enumerate(rows):
        style = styles["table_header"] if index == 0 else styles["table_body"]
        formatted_rows.append([_pdf_cell(value, style) for value in row])

    table = Table(formatted_rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#091A2B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#DFE6E9")),
                ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#F6F9F9")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
                    colors.white,
                    colors.HexColor("#F6F9F9"),
                ]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return table


def _pdf_summary_table(items, styles):
    return _pdf_table(
        [
            [item[0] for item in items],
            [item[1] for item in items],
        ],
        [540 / len(items)] * len(items),
        {
            **styles,
            "table_header": styles["summary_label"],
            "table_body": styles["summary_value"],
        },
    )


def _draw_pdf_header_footer(canvas, document):
    width, height = letter
    canvas.saveState()
    canvas.setFillColor(colors.HexColor("#091A2B"))
    canvas.rect(0, height - 36, width, 36, fill=1, stroke=0)
    canvas.setFillColor(colors.white)
    canvas.setFont(PDF_FONT_BOLD, 10)
    canvas.drawString(document.leftMargin, height - 23, "FinSight")
    canvas.setFont(PDF_FONT, 8)
    canvas.drawRightString(width - document.rightMargin, height - 23, "Financial Report")
    canvas.setFillColor(colors.HexColor("#71808D"))
    canvas.drawRightString(width - document.rightMargin, 20, f"Page {document.page}")
    canvas.restoreState()


def _build_pdf_report(report_data, currency="INR", language="en"):
    amount = lambda value: _pdf_amount(value, currency)
    label = lambda value: translate(value, language)
    styles = _pdf_styles()
    expenses = report_data.get("expenses") or {}
    expense_summary = expenses.get("expense_summary") or {}
    budgets = report_data.get("budgets") or {}
    investments = report_data.get("investments") or {}
    investment_stats = investments.get("stats") or {}
    goals = report_data.get("goals") or {}
    health = report_data.get("financial_health") or {}
    story = [
        Paragraph("FinSight", styles["title"]),
        Paragraph(label("Financial Report"), styles["subtitle"]),
        Paragraph(
            f"<b>{escape(label('Reporting Period'))}:</b> {escape(_pdf_period(report_data))}",
            styles["body"],
        ),
        Paragraph(
            f"<b>{escape(label('Generated'))}:</b> {escape(datetime.now().strftime('%B %d, %Y %H:%M'))}",
            styles["muted"],
        ),
    ]

    story.append(Paragraph(f"1. {label('Expense Summary')}", styles["section"]))
    story.append(
        _pdf_summary_table(
            [
                (label("Total Expenses"), amount(expense_summary.get("total_expenses"))),
                (label("Transaction Count"), _pdf_count(expense_summary.get("transaction_count"))),
                (label("Average Expense"), amount(expense_summary.get("average_expense"))),
                (label("Largest Expense"), amount(expense_summary.get("largest_expense"))),
            ],
            styles,
        )
    )

    story.append(Paragraph(f"2. {label('Expense by Category')}", styles["section"]))
    category_totals = expenses.get("category_totals") or []
    if category_totals:
        story.append(
            _pdf_table(
                [[label("Category"), label("Amount")]]
                + [
                    [item.get("category") or "Uncategorized", amount(item.get("amount"))]
                    for item in category_totals
                ],
                [360, 180],
                styles,
            )
        )
    else:
        story.append(Paragraph(label("No expense data available for the selected period."), styles["muted"]))

    story.append(Paragraph(f"3. {label('Monthly Spending')}", styles["section"]))
    monthly_totals = expenses.get("monthly_totals") or []
    if monthly_totals:
        story.append(
            _pdf_table(
                [[label("Month"), label("Amount")]]
                + [[item.get("month") or "Unknown", amount(item.get("amount"))] for item in monthly_totals],
                [360, 180],
                styles,
            )
        )
    else:
        story.append(Paragraph(label("No monthly expense data available."), styles["muted"]))

    story.append(Paragraph(f"4. {label('Budget Snapshot')} - {label('Current Snapshot')}", styles["section"]))
    story.append(Paragraph(label("Budget data represents the current stored budget snapshot."), styles["muted"]))
    story.append(
        _pdf_summary_table(
            [
                (label("Total Budget"), amount(budgets.get("total_budget"))),
                (label("Total Spent"), amount(budgets.get("total_spent"))),
                (label("Total Remaining"), amount(budgets.get("total_remaining"))),
                (label("Utilization"), _pdf_percent(budgets.get("overall_utilization"))),
                (label("Budget Count"), _pdf_count(budgets.get("budget_count"))),
            ],
            styles,
        )
    )
    budget_rows = budgets.get("budgets") or []
    if budget_rows:
        story.append(Spacer(1, 8))
        story.append(
            _pdf_table(
                [[label("Budget"), label("Category"), label("Budget Amount"), label("Spent"), label("Remaining"), label("Status")]]
                + [
                    [
                        item.get("budget_name") or "Budget",
                        item.get("category") or "Uncategorized",
                        amount(item.get("budget_amount")),
                        amount(item.get("spent_amount")),
                        amount(item.get("remaining_amount")),
                        label(item.get("status") or "Unavailable"),
                    ]
                    for item in budget_rows
                ],
                [120, 80, 90, 80, 90, 80],
                styles,
            )
        )

    story.append(Paragraph(f"5. {label('Investment Snapshot')} - {label('Current Snapshot')}", styles["section"]))
    story.append(Paragraph(label("Investment data represents the current portfolio snapshot."), styles["muted"]))
    story.append(
        _pdf_summary_table(
            [
                (label("Invested Value"), amount(investment_stats.get("total_invested"))),
                (label("Current Value"), amount(investment_stats.get("current_value"))),
                (label("Absolute Return"), amount(investment_stats.get("absolute_return"))),
                (label("Return Percentage"), _pdf_percent(investment_stats.get("return_percentage"))),
                (label("Holding Count"), _pdf_count(investments.get("holding_count"))),
            ],
            styles,
        )
    )
    allocation = investment_stats.get("allocation") or []
    if allocation:
        story.append(Spacer(1, 8))
        story.append(
            _pdf_table(
                [[label("Asset Type"), label("Allocation")]]
                + [[item.get("asset_type") or "Unknown", _pdf_percent(item.get("percentage"))] for item in allocation],
                [360, 180],
                styles,
            )
        )

    story.append(Paragraph(f"6. {label('Goal Progress')} - {label('Current Snapshot')}", styles["section"]))
    story.append(Paragraph(label("Goal values represent current saved progress; historical contributions are not available."), styles["muted"]))
    goal_rows = goals.get("goals") or []
    if goal_rows:
        story.append(
            _pdf_table(
                [[label("Goal"), label("Target"), label("Current"), label("Progress"), label("Remaining"), label("Status"), label("Target Date")]]
                + [
                    [
                        item.get("goal_name") or "Goal",
                        amount(item.get("target_amount")),
                        amount(item.get("current_amount")),
                        _pdf_percent(item.get("progress_percentage")),
                        amount(item.get("remaining_amount")),
                        label(item.get("status") or "Unavailable"),
                        _pdf_date(item.get("target_date")),
                    ]
                    for item in goal_rows
                ],
                [110, 75, 75, 65, 75, 80, 60],
                styles,
            )
        )
    else:
        story.append(Paragraph(label("No financial goals available."), styles["muted"]))

    story.append(Paragraph(f"7. {label('Financial Health')}", styles["section"]))
    if health.get("available"):
        story.append(
            _pdf_summary_table(
                [
                    (label("Score"), f"{_pdf_number(health.get('score'))}/100"),
                    (label("Grade"), label(health.get("grade") or "Unavailable")),
                    (label("Coverage"), _pdf_percent(health.get("coverage"))),
                    (label("Earned Points"), _pdf_number(health.get("earned_points"))),
                    (label("Available Weight"), _pdf_number(health.get("available_weight"))),
                ],
                styles,
            )
        )
        components = health.get("components") or {}
        if components:
            story.append(Spacer(1, 8))
            story.append(
                _pdf_table(
                    [[label("Component"), label("Status"), label("Score"), label("Details")]]
                    + [
                        [
                            name.replace("_", " ").title(),
                            label("Available" if component.get("available") else "Unavailable"),
                            _pdf_number(component.get("score")),
                            label(component.get("message") or ""),
                        ]
                        for name, component in components.items()
                    ],
                    [110, 80, 70, 280],
                    styles,
                )
            )
    else:
        story.append(Paragraph(label("Financial Health: Unavailable with the current data."), styles["muted"]))

    story.append(Paragraph(f"8. {label('Unavailable Metrics')}", styles["section"]))
    unavailable = report_data.get("unavailable_metrics") or {}
    unavailable_rows = [[label("Metric"), label("Status"), label("Reason")]]
    for name, metric in unavailable.items():
        if isinstance(metric, dict) and not metric.get("available", False):
            unavailable_rows.append(
                [
                    name.replace("_", " ").title(),
                    label("Unavailable"),
                    label(metric.get("reason") or "No reliable data source exists."),
                ]
            )
    if len(unavailable_rows) > 1:
        story.append(_pdf_table(unavailable_rows, [150, 80, 310], styles))
    else:
        story.append(Paragraph(label("No additional unavailable metrics were reported."), styles["muted"]))

    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=54,
        bottomMargin=34,
        title="FinSight Financial Report",
        author="FinSight",
        pageCompression=0,
    )
    document.build(story, onFirstPage=_draw_pdf_header_footer, onLaterPages=_draw_pdf_header_footer)
    return buffer.getvalue()


@app.get("/reports/export/pdf")
def export_pdf_report():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    try:
        report_data = build_reporting_data(
            current_user_id(),
            request.args.get("start_date"),
            request.args.get("end_date"),
            investment_service=investment_service,
            goal_service=goal_service,
        )
    except ReportValidationError as error:
        return jsonify({"error": str(error)}), 400

    preferences = get_user_preferences(current_user_id()) or PREFERENCE_DEFAULTS
    return Response(
        _build_pdf_report(
            report_data,
            preferences.get("currency", "USD"),
            preferences.get("language", "en"),
        ),
        mimetype="application/pdf",
        headers={
            "Content-Disposition": (
                "attachment; filename=\"finsight_financial_report.pdf\""
            )
        },
    )


EXCEL_PERCENT_FORMAT = "0.00%"


def _excel_currency_format(currency="INR"):
    decimals = 0 if str(currency or "").upper() == "JPY" else 2
    number_pattern = "#,##0" if decimals == 0 else f"#,##0.{('0' * decimals)}"
    return f'"{currency_symbol(currency)}" {number_pattern}'


def _excel_amount(value):
    number = _pdf_decimal(value)
    return "Unavailable" if number is None else float(number)


def _excel_percent(value):
    number = _pdf_decimal(value)
    return "Unavailable" if number is None else float(number) / 100


def _excel_number(value):
    number = _pdf_decimal(value)
    if number is None:
        return "Unavailable"
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def _excel_count(value):
    number = _pdf_decimal(value)
    return 0 if number is None else int(number)


def _excel_value(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def _excel_setup_sheet(workbook, name, title, report_data, language="en"):
    label = lambda text: translate(text, language)
    worksheet = workbook.create_sheet(name)
    worksheet.sheet_view.showGridLines = False
    worksheet.merge_cells("A1:H1")
    worksheet["A1"] = "FinSight"
    worksheet["A1"].font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    worksheet["A1"].fill = PatternFill("solid", fgColor="091A2B")
    worksheet["A1"].alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 26

    worksheet.merge_cells("A2:H2")
    worksheet["A2"] = label(title)
    worksheet["A2"].font = Font(name="Calibri", size=14, bold=True, color="138A70")
    worksheet["A2"].alignment = Alignment(vertical="center")
    worksheet.row_dimensions[2].height = 23

    worksheet["A3"] = label("Reporting Period")
    worksheet["B3"] = _pdf_period(report_data)
    worksheet["A4"] = label("Generated")
    worksheet["B4"] = datetime.now()
    worksheet["B4"].number_format = "mmmm d, yyyy h:mm AM/PM"
    for cell in (worksheet["A3"], worksheet["A4"]):
        cell.font = Font(bold=True, color="71808D")
    worksheet["B3"].alignment = Alignment(wrap_text=True)
    worksheet["B4"].alignment = Alignment(wrap_text=True)
    return worksheet


def _excel_section(worksheet, row, title, end_column):
    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=end_column,
    )
    cell = worksheet.cell(row=row, column=1, value=title)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="138A70")
    cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[row].height = 20


def _excel_table(worksheet, row, headers, rows, number_formats=None):
    number_formats = number_formats or {}
    header_fill = PatternFill("solid", fgColor="091A2B")
    border = Border(
        left=Side(style="thin", color="DFE6E9"),
        right=Side(style="thin", color="DFE6E9"),
        top=Side(style="thin", color="DFE6E9"),
        bottom=Side(style="thin", color="DFE6E9"),
    )

    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=row, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row_offset, values in enumerate(rows, start=1):
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(
                row=row + row_offset,
                column=column,
                value=_excel_value(value),
            )
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if column in number_formats:
                cell.number_format = number_formats[column]

    return row + len(rows) + 1


def _excel_empty(worksheet, row, message, end_column):
    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=end_column,
    )
    cell = worksheet.cell(row=row, column=1, value=message)
    cell.font = Font(italic=True, color="71808D")
    cell.alignment = Alignment(wrap_text=True)


def _excel_set_widths(worksheet, widths):
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column)].width = width


def _excel_finalize(worksheet, freeze_panes):
    worksheet.freeze_panes = freeze_panes
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_setup.orientation = "landscape"


def _build_excel_report(report_data, currency="INR", language="en"):
    excel_currency_format = _excel_currency_format(currency)
    label = lambda text: translate(text, language)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "FinSight Financial Report"
    workbook.properties.creator = "FinSight"

    expenses = report_data.get("expenses") or {}
    expense_summary = expenses.get("expense_summary") or {}
    budgets = report_data.get("budgets") or {}
    investments = report_data.get("investments") or {}
    investment_stats = investments.get("stats") or {}
    goals = report_data.get("goals") or {}
    health = report_data.get("financial_health") or {}

    summary = _excel_setup_sheet(workbook, "Summary", "Financial Report", report_data, language)
    _excel_section(summary, 6, label("Expense Summary"), 4)
    _excel_table(
        summary,
        7,
        [label("Total Expenses"), label("Transaction Count"), label("Average Expense"), label("Largest Expense")],
        [[
            _excel_amount(expense_summary.get("total_expenses")),
            _excel_count(expense_summary.get("transaction_count")),
            _excel_amount(expense_summary.get("average_expense")),
            _excel_amount(expense_summary.get("largest_expense")),
        ]],
        {1: excel_currency_format, 3: excel_currency_format, 4: excel_currency_format},
    )
    _excel_section(summary, 10, label("Financial Health"), 3)
    _excel_table(
        summary,
        11,
        [label("Overall Score"), label("Grade"), label("Coverage")],
        [[
            _excel_number(health.get("score")),
            label(health.get("grade") or "Unavailable"),
            _excel_percent(health.get("coverage")),
        ]],
        {3: EXCEL_PERCENT_FORMAT},
    )
    unavailable_rows = []
    for name, metric in (report_data.get("unavailable_metrics") or {}).items():
        if isinstance(metric, dict) and not metric.get("available", False):
            unavailable_rows.append([
                name.replace("_", " ").title(),
                label("Unavailable"),
                label(metric.get("reason") or "No reliable data source exists."),
            ])
    _excel_section(summary, 14, label("Unavailable Metrics"), 3)
    if unavailable_rows:
        _excel_table(summary, 15, [label("Metric"), label("Status"), label("Reason")], unavailable_rows)
    else:
        _excel_empty(summary, 15, label("No additional unavailable metrics were reported."), 3)
    _excel_set_widths(summary, [24, 32, 24, 24, 14, 14, 14, 14])
    _excel_finalize(summary, None)

    expense_sheet = _excel_setup_sheet(workbook, "Expense Analysis", "Expense Analysis", report_data, language)
    _excel_section(expense_sheet, 6, label("Expense Summary"), 4)
    _excel_table(
        expense_sheet,
        7,
        [label("Total Expenses"), label("Transaction Count"), label("Average Expense"), label("Largest Expense")],
        [[
            _excel_amount(expense_summary.get("total_expenses")),
            _excel_count(expense_summary.get("transaction_count")),
            _excel_amount(expense_summary.get("average_expense")),
            _excel_amount(expense_summary.get("largest_expense")),
        ]],
        {1: excel_currency_format, 3: excel_currency_format, 4: excel_currency_format},
    )
    _excel_section(expense_sheet, 10, label("Category Analysis"), 2)
    category_rows = [
        [item.get("category") or "Uncategorized", _excel_amount(item.get("amount"))]
        for item in (expenses.get("category_totals") or [])
    ]
    if category_rows:
        _excel_table(expense_sheet, 11, [label("Category"), label("Amount")], category_rows, {2: excel_currency_format})
    else:
        _excel_empty(expense_sheet, 11, label("No expense data available for the selected period."), 2)
    monthly_start = 12 + len(category_rows)
    _excel_section(expense_sheet, monthly_start, label("Monthly Analysis"), 2)
    monthly_rows = [
        [item.get("month") or "Unknown", _excel_amount(item.get("amount"))]
        for item in (expenses.get("monthly_totals") or [])
    ]
    if monthly_rows:
        _excel_table(expense_sheet, monthly_start + 1, [label("Month"), label("Amount")], monthly_rows, {2: excel_currency_format})
    else:
        _excel_empty(expense_sheet, monthly_start + 1, label("No monthly expense data available."), 2)
    _excel_set_widths(expense_sheet, [30, 22, 24, 24, 14, 14, 14, 14])
    _excel_finalize(expense_sheet, "A8")

    budget_sheet = _excel_setup_sheet(workbook, "Budget Summary", "Budget Summary - Current Snapshot", report_data, language)
    _excel_section(budget_sheet, 6, label("Current Budget Snapshot"), 5)
    _excel_table(
        budget_sheet,
        7,
        [label("Total Budget"), label("Total Spent"), label("Total Remaining"), label("Overall Utilization"), label("Budget Count")],
        [[
            _excel_amount(budgets.get("total_budget")),
            _excel_amount(budgets.get("total_spent")),
            _excel_amount(budgets.get("total_remaining")),
            _excel_percent(budgets.get("overall_utilization")),
            _excel_count(budgets.get("budget_count")),
        ]],
        {
            1: excel_currency_format,
            2: excel_currency_format,
            3: excel_currency_format,
            4: EXCEL_PERCENT_FORMAT,
        },
    )
    _excel_empty(budget_sheet, 10, label("Budget data represents the current stored budget snapshot."), 6)
    _excel_section(budget_sheet, 12, label("Budgets"), 6)
    budget_rows = [
        [
            item.get("budget_name") or "Budget",
            item.get("category") or "Uncategorized",
            _excel_amount(item.get("budget_amount")),
            _excel_amount(item.get("spent_amount")),
            _excel_amount(item.get("remaining_amount")),
            label(item.get("status") or "Unavailable"),
        ]
        for item in (budgets.get("budgets") or [])
    ]
    if budget_rows:
        _excel_table(
            budget_sheet,
            13,
            [label("Budget Name"), label("Category"), label("Budget Amount"), label("Spent Amount"), label("Remaining Amount"), label("Status")],
            budget_rows,
            {3: excel_currency_format, 4: excel_currency_format, 5: excel_currency_format},
        )
    else:
        _excel_empty(budget_sheet, 13, label("No current budget data available."), 6)
    _excel_set_widths(budget_sheet, [26, 20, 18, 18, 20, 18, 14, 14])
    _excel_finalize(budget_sheet, "A8")

    investment_sheet = _excel_setup_sheet(workbook, "Investments", "Investments - Current Snapshot", report_data, language)
    _excel_section(investment_sheet, 6, label("Current Portfolio Snapshot"), 5)
    _excel_table(
        investment_sheet,
        7,
        [label("Invested Value"), label("Current Value"), label("Absolute Return"), label("Return Percentage"), label("Holding Count")],
        [[
            _excel_amount(investment_stats.get("total_invested")),
            _excel_amount(investment_stats.get("current_value")),
            _excel_amount(investment_stats.get("absolute_return")),
            _excel_percent(investment_stats.get("return_percentage")),
            _excel_count(investments.get("holding_count")),
        ]],
        {
            1: excel_currency_format,
            2: excel_currency_format,
            3: excel_currency_format,
            4: EXCEL_PERCENT_FORMAT,
        },
    )
    _excel_empty(investment_sheet, 10, label("Investment data represents the current portfolio snapshot."), 7)
    _excel_section(investment_sheet, 12, label("Holdings"), 7)
    holding_rows = [
        [
            item.get("asset_name") or "Asset",
            item.get("asset_type") or "Unknown",
            item.get("quantity"),
            _excel_amount(item.get("purchase_price")),
            _excel_amount(item.get("current_value")),
            _excel_amount(item.get("invested_value")),
            item.get("purchase_date") or label("Unavailable"),
        ]
        for item in (investments.get("holdings") or [])
    ]
    if holding_rows:
        _excel_table(
            investment_sheet,
            13,
            [label("Asset Name"), label("Asset Type"), label("Quantity"), label("Purchase Price"), label("Current Value"), label("Invested Value"), label("Purchase Date")],
            holding_rows,
            {4: excel_currency_format, 5: excel_currency_format, 6: excel_currency_format},
        )
    else:
        _excel_empty(investment_sheet, 13, label("No investment holdings available."), 7)
    allocation_start = 14 + len(holding_rows)
    _excel_section(investment_sheet, allocation_start, label("Asset Allocation"), 2)
    allocation_rows = [
        [item.get("asset_type") or "Unknown", _excel_percent(item.get("percentage"))]
        for item in (investment_stats.get("allocation") or [])
    ]
    if allocation_rows:
        _excel_table(investment_sheet, allocation_start + 1, [label("Asset Type"), label("Allocation")], allocation_rows, {2: EXCEL_PERCENT_FORMAT})
    else:
        _excel_empty(investment_sheet, allocation_start + 1, label("No asset allocation data available."), 2)
    _excel_set_widths(investment_sheet, [25, 18, 14, 18, 18, 18, 18, 14])
    _excel_finalize(investment_sheet, "A8")

    goal_sheet = _excel_setup_sheet(workbook, "Goals", "Goals - Current Progress", report_data, language)
    _excel_section(goal_sheet, 6, label("Current Goal Progress"), 7)
    goal_rows = [
        [
            item.get("goal_name") or "Goal",
            item.get("goal_category") or "Uncategorized",
            _excel_amount(item.get("target_amount")),
            _excel_amount(item.get("current_amount")),
            _excel_percent(item.get("progress_percentage")),
            _excel_amount(item.get("remaining_amount")),
            label(item.get("status") or "Unavailable"),
        ]
        for item in (goals.get("goals") or [])
    ]
    if goal_rows:
        _excel_table(
            goal_sheet,
            7,
            [label("Goal Name"), label("Category"), label("Target Amount"), label("Current Amount"), label("Progress Percentage"), label("Remaining Amount"), label("Status")],
            goal_rows,
            {
                3: excel_currency_format,
                4: excel_currency_format,
                5: EXCEL_PERCENT_FORMAT,
                6: excel_currency_format,
            },
        )
    else:
        _excel_empty(goal_sheet, 7, label("No financial goals available."), 7)
    _excel_set_widths(goal_sheet, [25, 18, 18, 18, 20, 20, 18, 14])
    _excel_finalize(goal_sheet, "A8")

    health_sheet = _excel_setup_sheet(workbook, "Financial Health", "Financial Health", report_data, language)
    _excel_section(health_sheet, 6, label("Existing Financial Health Score"), 5)
    if health.get("available"):
        _excel_table(
            health_sheet,
            7,
            [label("Overall Score"), label("Grade"), label("Coverage"), label("Earned Points"), label("Available Weight")],
            [[
                _excel_number(health.get("score")),
                label(health.get("grade") or "Unavailable"),
                _excel_percent(health.get("coverage")),
                _excel_number(health.get("earned_points")),
                _excel_number(health.get("available_weight")),
            ]],
            {3: EXCEL_PERCENT_FORMAT},
        )
        component_rows = [
            [
                name.replace("_", " ").title(),
                label("Available" if component.get("available") else "Unavailable"),
                _excel_number(component.get("score")),
                label(component.get("message") or ""),
            ]
            for name, component in (health.get("components") or {}).items()
        ]
        _excel_section(health_sheet, 10, label("Components"), 4)
        if component_rows:
            _excel_table(health_sheet, 11, [label("Component"), label("Status"), label("Score"), label("Message")], component_rows)
        else:
            _excel_empty(health_sheet, 11, label("No Financial Health component details available."), 4)
    else:
        _excel_empty(health_sheet, 7, label("Financial Health is unavailable with the current data."), 5)
    _excel_set_widths(health_sheet, [24, 18, 16, 70, 14, 14, 14, 14])
    _excel_finalize(health_sheet, "A8")

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@app.get("/reports/export/excel")
def export_excel_report():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    try:
        report_data = build_reporting_data(
            current_user_id(),
            request.args.get("start_date"),
            request.args.get("end_date"),
            investment_service=investment_service,
            goal_service=goal_service,
        )
    except ReportValidationError as error:
        return jsonify({"error": str(error)}), 400

    preferences = get_user_preferences(current_user_id()) or PREFERENCE_DEFAULTS
    return Response(
        _build_excel_report(
            report_data,
            preferences.get("currency", "USD"),
            preferences.get("language", "en"),
        ),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                "attachment; filename=\"finsight_financial_report.xlsx\""
            )
        },
    )


@app.route("/expense")
def expense():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    stats = get_summary_stats(current_user_id())
    return render_template(
        "expenses/dashboard.html",
        stats=stats,
        current_username=session["username"],
    )

@app.route("/budget")
@app.route("/budgets")
def budgets():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    search_query = request.args.get("q", "").strip()
    category_filter = request.args.get("category", "All")
    status_filter = request.args.get("status", "All")
    priority_filter = request.args.get("priority", "All")
    sort_by = request.args.get("sort", "newest")

    budget_rows = filter_budgets(
        current_user_id(),
        search_query=search_query,
        category_filter=category_filter,
        status_filter=status_filter,
        priority_filter=priority_filter,
        sort_by=sort_by,
    )
    stats = get_summary_stats(current_user_id())

    return render_template(
        "budgets/dashboard.html",
        budgets=budget_rows,
        stats=stats,
        search_query=search_query,
        category_filter=category_filter,
        status_filter=status_filter,
        priority_filter=priority_filter,
        sort_by=sort_by,
        current_username=session["username"],
        budget_csrf_token=budget_csrf_token(),
    )
@app.context_processor
def inject_template_globals():
    return {
        "today_date": datetime.now().strftime("%B %d, %Y"),
        "current_username": session.get("username", ""),
    }


def validate_budget_form(form):
    errors = []
    budget_name = form.get("budget_name", "").strip()
    category = form.get("category", "").strip()
    start_date = form.get("start_date", "").strip()
    end_date = form.get("end_date", "").strip()

    if not budget_name:
        errors.append("Budget name is required.")
    if not category:
        errors.append("Category is required.")

    try:
        if float(form.get("budget_amount", 0) or 0) < 0:
            errors.append("Budget amount cannot be negative.")
    except ValueError:
        errors.append("Budget amount must be a valid number.")

    try:
        if float(form.get("spent_amount", 0) or 0) < 0:
            errors.append("Spent amount cannot be negative.")
    except ValueError:
        errors.append("Spent amount must be a valid number.")

    if not start_date or not end_date:
        errors.append("Start date and end date are required.")
    elif start_date > end_date:
        errors.append("Start date cannot be later than end date.")

    return errors


def render_budget_form(budget=None, is_edit=False, budget_id=None):
    return render_template(
        "budgets/form.html",
        budget=budget or {},
        is_edit=is_edit,
        budget_id=budget_id,
        budget_csrf_token=budget_csrf_token(),
    )

@app.route("/budget/create", methods=["GET", "POST"])
def create_budget_route():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    if request.method == "POST":
        if not budget_csrf_valid():
            flash("The form security token is missing or invalid.", "danger")
            return render_budget_form(request.form, is_edit=False), 400
        errors = validate_budget_form(request.form)
        if errors:
            for error in errors:
                flash(error, "danger")
            return render_budget_form(request.form, is_edit=False)

        try:
            create_budget(current_user_id(), request.form)
            flash("Budget created successfully.", "success")
            return redirect(url_for("budgets"))
        except Exception:
            app.logger.exception("Budget creation failed")
            flash(
                "Unable to create budget. Please check the form and try again.",
                "danger",
            )
            return render_budget_form(request.form, is_edit=False)

    return render_budget_form(is_edit=False)


@app.route("/budget/edit/<int:budget_id>", methods=["GET", "POST"])
def edit_budget(budget_id):
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    budget = get_budget(budget_id, current_user_id())
    if not budget:
        flash("Budget not found.", "danger")
        return redirect(url_for("budgets"))

    if request.method == "POST":
        if not budget_csrf_valid():
            flash("The form security token is missing or invalid.", "danger")
            return render_budget_form(request.form, is_edit=True, budget_id=budget_id), 400
        errors = validate_budget_form(request.form)
        if errors:
            for error in errors:
                flash(error, "danger")
            return render_budget_form(request.form, is_edit=True, budget_id=budget_id)

        try:
            if update_budget(budget_id, current_user_id(), request.form):
                flash("Budget updated successfully.", "success")
                return redirect(url_for("budgets"))
            flash("Unable to update budget.", "danger")
        except Exception:
            app.logger.exception("Budget update failed")
            flash(
                "Unable to update budget. Please check the form and try again.",
                "danger",
            )

    return render_budget_form(budget=budget, is_edit=True, budget_id=budget_id)


@app.route("/budget/view/<int:budget_id>")
def view_budget(budget_id):
    if login_required_redirect():
        return jsonify({"success": False, "message": "Unauthorized"}), 401

    budget = get_budget(budget_id, current_user_id())
    if not budget:
        return jsonify({"success": False, "message": "Budget not found"}), 404

    return jsonify({"success": True, "budget": budget})


@app.post("/budget/delete/<int:budget_id>")
def remove_budget(budget_id):
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response
    if not budget_csrf_valid():
        flash("The form security token is missing or invalid.", "danger")
        return redirect(url_for("budgets"))

    result = delete_budget(budget_id, current_user_id())

    if result:
        flash("Budget deleted successfully.", "success")
    else:
        flash("Budget not found.", "danger")

    return redirect(url_for("budgets"))


EXPENSE_CATEGORIES = [
    "Food & Dining",
    "Groceries",
    "Transportation",
    "Utilities",
    "Housing & Rent",
    "Healthcare",
    "Shopping",
    "Entertainment",
    "Travel",
    "Personal Care",
    "Education",
    "Subscriptions",
    "Other",
]

PAYMENT_MODES = [
    "Cash",
    "Credit Card",
    "Debit Card",
    "UPI",
    "Bank Transfer",
    "Wallet",
]


def validate_expense_form(form):
    errors = []
    category = form.get("category", "").strip()
    payment_mode = form.get("payment_mode", "").strip()
    expense_date = form.get("date", "").strip()

    try:
        amount = float(form.get("amount", 0) or 0)
        if amount <= 0:
            errors.append("Amount must be greater than zero.")
    except ValueError:
        errors.append("Amount must be a valid number.")

    if category not in EXPENSE_CATEGORIES:
        errors.append("Please select a valid category.")

    if payment_mode not in PAYMENT_MODES:
        errors.append("Please select a valid payment mode.")

    if not expense_date:
        errors.append("Date is required.")
    else:
        try:
            parsed_date = datetime.strptime(expense_date, "%Y-%m-%d").date()
            if parsed_date > datetime.now().date():
                errors.append("Date cannot be in the future.")
        except ValueError:
            errors.append("Date must be valid.")

    return errors


def render_expense_form(expense_row=None, is_edit=False, transaction_id=None):
    return render_template(
        "expenses/form.html",
        expense=expense_row or {},
        is_edit=is_edit,
        transaction_id=transaction_id,
        categories=EXPENSE_CATEGORIES,
        payment_modes=PAYMENT_MODES,
        expense_csrf_token=expense_csrf_token(),
    )


@app.route("/expenses")
def expenses():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    from db import get_expense_summary, get_transactions

    search_query = request.args.get("q", "").strip()
    category_filter = request.args.get("category", "All")
    payment_filter = request.args.get("payment", "All")
    sort_by = request.args.get("sort", "newest")

    expense_rows = get_transactions(
        current_user_id(),
        search_query=search_query,
        category_filter=category_filter,
        payment_filter=payment_filter,
        sort_by=sort_by,
    )
    stats = get_expense_summary(current_user_id())

    return render_template(
        "expenses/dashboard.html",
        expenses=expense_rows,
        stats=stats,
        categories=EXPENSE_CATEGORIES,
        payment_modes=PAYMENT_MODES,
        search_query=search_query,
        category_filter=category_filter,
        payment_filter=payment_filter,
        sort_by=sort_by,
        current_username=session["username"],
        expense_csrf_token=expense_csrf_token(),
    )


CSV_FORMULA_PREFIXES = ("=", "+", "-", "@")


def _csv_safe_text(value):
    if isinstance(value, str) and value.startswith(CSV_FORMULA_PREFIXES):
        return "'" + value
    return value


@app.get("/expenses/export")
def export_expenses():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    expense_rows = get_transactions(current_user_id())
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["Date", "Category", "Description", "Payment Mode", "Type", "Amount"])

    for expense in expense_rows:
        writer.writerow(
            [
                _csv_safe_text(expense.get("date", "")),
                _csv_safe_text(expense.get("category", "")),
                _csv_safe_text(expense.get("description") or ""),
                _csv_safe_text(expense.get("payment_mode", "")),
                _csv_safe_text(expense.get("type", "")),
                expense.get("amount", ""),
            ]
        )

    response = Response(output.getvalue(), mimetype="text/csv")
    response.headers["Content-Disposition"] = (
        "attachment; filename=finsight_expenses.csv"
    )
    return response


@app.route("/expense/create", methods=["GET", "POST"])
def create_expense():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    if request.method == "POST":
        if not expense_csrf_valid():
            flash("The form security token is missing or invalid.", "danger")
            return render_expense_form(request.form, is_edit=False), 400
        errors = validate_expense_form(request.form)
        if errors:
            for error in errors:
                flash(error, "danger")
            return render_expense_form(request.form, is_edit=False)

        try:
            from db import create_transaction

            create_transaction(current_user_id(), request.form)
            flash("Expense added successfully.", "success")
            return redirect(url_for("expenses"))
        except Exception:
            app.logger.exception("Expense creation failed")
            flash("Unable to add expense. Please check the form and try again.", "danger")
            return render_expense_form(request.form, is_edit=False)

    return render_expense_form(is_edit=False)


@app.route("/expense/edit/<int:transaction_id>", methods=["GET", "POST"])
def edit_expense(transaction_id):
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    from db import get_transaction

    expense_row = get_transaction(transaction_id, current_user_id())
    if not expense_row:
        flash("Expense not found.", "danger")
        return redirect(url_for("expenses"))

    if request.method == "POST":
        if not expense_csrf_valid():
            flash("The form security token is missing or invalid.", "danger")
            return render_expense_form(
                request.form,
                is_edit=True,
                transaction_id=transaction_id,
            ), 400
        errors = validate_expense_form(request.form)
        if errors:
            for error in errors:
                flash(error, "danger")
            return render_expense_form(
                request.form,
                is_edit=True,
                transaction_id=transaction_id,
            )

        try:
            from db import update_transaction

            if update_transaction(transaction_id, current_user_id(), request.form):
                flash("Expense updated successfully.", "success")
                return redirect(url_for("expenses"))
            flash("Unable to update expense.", "danger")
        except Exception:
            app.logger.exception("Expense update failed")
            flash("Unable to update expense. Please check the form and try again.", "danger")

    return render_expense_form(
        expense_row=expense_row,
        is_edit=True,
        transaction_id=transaction_id,
    )


@app.route("/expense/view/<int:transaction_id>")
def view_expense(transaction_id):
    if login_required_redirect():
        return jsonify({"success": False, "message": "Unauthorized"}), 401

    from db import get_transaction

    expense_row = get_transaction(transaction_id, current_user_id())
    if not expense_row:
        return jsonify({"success": False, "message": "Expense not found"}), 404

    return jsonify({"success": True, "expense": expense_row})


@app.post("/expense/delete/<int:transaction_id>")
def delete_expense_route(transaction_id):
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response
    if not expense_csrf_valid():
        flash("The form security token is missing or invalid.", "danger")
        return redirect(url_for("expenses"))

    from db import delete_transaction

    result = delete_transaction(transaction_id, current_user_id())
    if result:
        flash("Expense deleted successfully.", "success")
    else:
        flash("Expense not found.", "danger")

    return redirect(url_for("expenses"))


def render_investment_form(investment=None, is_edit=False, investment_id=None, errors=None):
    return render_template(
        "investments/form.html",
        investment=investment or {},
        is_edit=is_edit,
        investment_id=investment_id,
        asset_types=ASSET_TYPES,
        investment_errors=errors or [],
        investment_csrf_token=investment_csrf_token(),
    )


@app.route("/investments")
def investments():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    user_id = current_user_id()
    rows, stats = investment_service.list_for_user(user_id)
    goals = goal_service.list_for_user(user_id)
    return render_template(
        "investments/dashboard.html",
        investments=rows,
        stats=stats,
        goals=goals,
    )


@app.route("/investments/create", methods=["GET", "POST"])
def create_investment():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    if request.method == "GET":
        return render_investment_form()
    if not investment_csrf_valid():
        flash("The form security token is missing or invalid.", "danger")
        return render_investment_form(request.form, errors=["Invalid form security token."]), 400

    try:
        created, errors = investment_service.create(current_user_id(), request.form)
    except Exception:
        app.logger.exception("Investment creation failed")
        flash("Unable to create investment. Please try again.", "danger")
        return render_investment_form(request.form), 503
    if errors:
        for error in errors:
            flash(error, "danger")
        return render_investment_form(request.form, errors=errors), 400

    flash("Investment added successfully.", "success")
    return redirect(url_for("investments"))


@app.route("/investments/<int:investment_id>")
def view_investment(investment_id):
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    investment = investment_service.get_for_user(investment_id, current_user_id())
    if not investment:
        flash("Investment not found.", "danger")
        return redirect(url_for("investments"))
    return render_template("investments/detail.html", investment=investment)


@app.route("/investments/<int:investment_id>/edit", methods=["GET", "POST"])
def edit_investment(investment_id):
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response

    investment = investment_service.get_for_user(investment_id, current_user_id())
    if not investment:
        flash("Investment not found.", "danger")
        return redirect(url_for("investments"))
    if request.method == "GET":
        return render_investment_form(investment, True, investment_id)
    if not investment_csrf_valid():
        flash("The form security token is missing or invalid.", "danger")
        return render_investment_form(request.form, True, investment_id, ["Invalid form security token."]), 400

    try:
        updated, errors = investment_service.update(investment_id, current_user_id(), request.form)
    except Exception:
        app.logger.exception("Investment update failed")
        flash("Unable to update investment. Please try again.", "danger")
        return render_investment_form(request.form, True, investment_id), 503
    if errors:
        for error in errors:
            flash(error, "danger")
        return render_investment_form(request.form, True, investment_id, errors), 400
    if not updated:
        flash("Investment not found.", "danger")
        return redirect(url_for("investments"))

    flash("Investment updated successfully.", "success")
    return redirect(url_for("investments"))


@app.post("/investments/<int:investment_id>/delete")
def delete_investment(investment_id):
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response
    if not investment_csrf_valid():
        flash("The form security token is missing or invalid.", "danger")
        return redirect(url_for("investments"))

    deleted = investment_service.delete(investment_id, current_user_id())
    flash("Investment deleted successfully." if deleted else "Investment not found.",
          "success" if deleted else "danger")
    return redirect(url_for("investments"))


def render_goal_form(goal=None, is_edit=False, goal_id=None, errors=None):
    return render_template(
        "goals/form.html",
        goal=goal or {},
        is_edit=is_edit,
        goal_id=goal_id,
        goal_categories=GOAL_CATEGORIES,
        goal_errors=errors or [],
        goal_csrf_token=goal_csrf_token(),
    )


@app.route("/goals")
def goals():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response
    return render_template("goals/dashboard.html", goals=goal_service.list_for_user(current_user_id()))


@app.route("/goals/create", methods=["GET", "POST"])
def create_goal():
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response
    if request.method == "GET":
        return render_goal_form()
    if not goal_csrf_valid():
        return render_goal_form(request.form, errors=["Invalid form security token."]), 400

    created, errors = goal_service.create(current_user_id(), request.form)
    if errors:
        for error in errors:
            flash(error, "danger")
        return render_goal_form(request.form, errors=errors), 400
    flash("Financial goal created successfully.", "success")
    return redirect(url_for("goals"))


@app.route("/goals/<int:goal_id>")
def view_goal(goal_id):
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response
    goal = goal_service.get_for_user(goal_id, current_user_id())
    if not goal:
        flash("Goal not found.", "danger")
        return redirect(url_for("goals"))
    return render_template("goals/detail.html", goal=goal)


@app.route("/goals/<int:goal_id>/edit", methods=["GET", "POST"])
def edit_goal(goal_id):
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response
    goal = goal_service.get_for_user(goal_id, current_user_id())
    if not goal:
        flash("Goal not found.", "danger")
        return redirect(url_for("goals"))
    if request.method == "GET":
        return render_goal_form(goal, True, goal_id)
    if not goal_csrf_valid():
        return render_goal_form(request.form, True, goal_id, ["Invalid form security token."]), 400

    updated, errors = goal_service.update(goal_id, current_user_id(), request.form)
    if errors:
        for error in errors:
            flash(error, "danger")
        return render_goal_form(request.form, True, goal_id, errors), 400
    if not updated:
        flash("Goal not found.", "danger")
        return redirect(url_for("goals"))
    flash("Financial goal updated successfully.", "success")
    return redirect(url_for("goals"))


@app.post("/goals/<int:goal_id>/delete")
def delete_goal(goal_id):
    redirect_response = login_required_redirect()
    if redirect_response:
        return redirect_response
    if not goal_csrf_valid():
        flash("The form security token is missing or invalid.", "danger")
        return redirect(url_for("goals"))
    deleted = goal_service.delete(goal_id, current_user_id())
    flash("Goal deleted successfully." if deleted else "Goal not found.",
          "success" if deleted else "danger")
    return redirect(url_for("goals"))


if __name__ == "__main__":
    debug = _debug_enabled()
    app.run(debug=debug, use_reloader=debug)
